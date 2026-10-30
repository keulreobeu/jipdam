import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl

from budongi.team_collection import contract as C
from budongi.team_collection.xcheck.xlsx_cells import iter_sheet_rows, sheet_names, split_ref

NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
RNS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
PNS = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
WS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
SS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"


def build(path, sheets, shared=None):
    wb = "".join(f'<sheet name="{n}" sheetId="{i + 1}" r:id="rId{i + 1}"/>' for i, (n, _) in enumerate(sheets))
    rels = "".join(f'<Relationship Id="rId{i + 1}" Type="{WS}" Target="worksheets/sheet{i + 1}.xml"/>'
                   for i in range(len(sheets)))
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("xl/workbook.xml", f"<workbook {NS} {RNS}><sheets>{wb}</sheets></workbook>")
        if shared is not None:
            rels += f'<Relationship Id="rSS" Type="{SS}" Target="sharedStrings.xml"/>'
            z.writestr("xl/sharedStrings.xml", f"<sst {NS}>" + "".join(shared) + "</sst>")
        z.writestr("xl/_rels/workbook.xml.rels", f"<Relationships {PNS}>{rels}</Relationships>")
        for i, (_, data) in enumerate(sheets):
            z.writestr(f"xl/worksheets/sheet{i + 1}.xml", f"<worksheet {NS}><sheetData>{data}</sheetData></worksheet>")
    return path


class SplitRefTests(unittest.TestCase):
    def test_simple_and_multi_letter(self):
        self.assertEqual(split_ref("A1"), (1, 1))
        self.assertEqual(split_ref("AB12"), (28, 12))

    def test_lowercase_rejected(self):
        with self.assertRaises(C.ContractError) as cm:
            split_ref("a1")
        self.assertEqual(str(cm.exception), "invalid_xlsx:cell_ref")


class SheetReadTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_openpyxl_workbook(self):
        p = self.dir / "o.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "데이터"
        ws.append(["name", "code", None, "n"])
        ws.append(["서울", "00123", None, 12])
        ws["D3"] = 1.5
        ws["E3"] = True
        ws["F3"] = "=D3+1"
        wb.create_sheet("Second")["B2"] = "x"
        wb.save(p)
        self.assertEqual(sheet_names(p), ["데이터", "Second"])
        rows = list(iter_sheet_rows(p, "데이터"))
        self.assertEqual(rows, [
            (1, ("name", "code", None, "n")),
            (2, ("서울", "00123", None, "12")),
            (3, (None, None, None, "1.5", "TRUE")),
        ])
        self.assertEqual(list(iter_sheet_rows(str(p), "Second")), [(2, (None, "x"))])

    def test_shared_inline_str_bool(self):
        p = build(self.dir / "h.xlsx", [("S", (
            '<row r="1"><c r="A1" t="s"><v>1</v></c><c r="B1" t="s"><v>0</v></c></row>'
            '<row r="2"><c r="A2" t="inlineStr"><is><t>x</t></is></c><c r="C2"><v>7</v></c></row>'
            '<row r="3"><c r="B3" t="b"><v>0</v></c><c r="C3" t="str"><v>ok</v></c></row>'))],
            shared=["<si><t>alpha</t></si>", "<si><t>beta</t></si>"])
        self.assertEqual(list(iter_sheet_rows(p, "S")), [
            (1, ("beta", "alpha")),
            (2, ("x", None, "7")),
            (3, (None, "FALSE", "ok")),
        ])

    def test_unknown_sheet(self):
        p = build(self.dir / "u.xlsx", [("S", "")])
        with self.assertRaises(C.ContractError) as cm:
            list(iter_sheet_rows(p, "Missing"))
        self.assertEqual(str(cm.exception), "unknown_sheet")


if __name__ == "__main__":
    unittest.main()
