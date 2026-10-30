import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from budongi.team_collection.base import Cell
from budongi.team_collection.contract import ContractError
from budongi.team_collection.workbook.writer import write_workbook

NAMES = ("데이터셋 목록", "컬럼별 메타데이터", "데이터 품질 및 연계", "전처리 및 이슈 관리")
_H = (
    "데이터셋_ID, 데이터셋명, 분류, 담당자, 제공기관, 출처_URL, 수집방법, 전체행수, 전체컬럼수, 파일수, 파일형식, 인코딩, 데이터기준일, 수집일, 갱신주기, 공간범위, 행단위, PK후보, JOIN키, 좌표계, 저장위치, 가공상태, 이용조건, 활용목적, 중복관계, 측정근거, 검증상태, 원본해시, 수집기록",
    "데이터셋_ID, 컬럼명, 한글명, 설명, 물리자료형, 논리자료형, 전체행수, NULL수, NULL비율, 고유값수, 고유값비율, 최솟값, 최댓값, 평균, 중앙값, 대표값, 단위, 허용값_코드, 이상값수, 이상값비율, 키구분, 전처리필요여부, 의미상결측수, 통합결측수, 통합결측률, 시트_원본위치, 검사대상수, 유효성검사규칙, 측정근거, 정리본_열",
    "데이터셋_ID, 전체중복건수, 전체중복률, PK중복건수, PK중복률, 유효성오류건수, 주요키결측률, 주소보유율, 좌표보유율, 좌표유효율, 공간정보형식, 좌표계, 행정구역기준, 연결대상, JOIN키_방식, JOIN성공률, 미연결건수, 기준시점일치여부, 전체행수, 중복판정기준, 키기준, 주요키결측수, 주소보유수, 좌표보유수, 좌표유효수, 검사대상수, 검사규칙, 검증상태, 근거",
    "이슈_ID, 데이터셋_ID, 대상컬럼, 문제유형, 상세내용, 영향도, 처리방안, 담당자, 처리상태, 처리일, 비고",
)
HEADERS = {n: tuple(h.split(", ")) for n, h in zip(NAMES, _H)}
NR = 'IF((G6)=0,"해당 없음",H6/(G6))'


def make_sheets():
    sheets = {n: {"title": "T " + n, "note": "N " + n, "headers": HEADERS[n], "rows": []} for n in NAMES}
    col = [Cell(None) for _ in HEADERS[NAMES[1]]]
    col[0] = Cell("T001")
    col[6] = Cell(100, None, "#,##0")
    col[8] = Cell(None, NR, "0.00%")
    sheets[NAMES[1]]["rows"] = [col]
    sheets[NAMES[3]]["rows"] = [[Cell("T-I%03d" % i)] + [Cell(None)] * 10 for i in (1, 2)]
    return sheets


class WriterTypicalTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "meta.xlsx"
        self.result = write_workbook(self.path, sheets=make_sheets())
        self.wb = load_workbook(self.path)

    def tearDown(self):
        self._tmp.cleanup()

    def test_return_and_sheet_order(self):
        self.assertEqual(self.result, {NAMES[0]: 0, NAMES[1]: 1, NAMES[2]: 0, NAMES[3]: 2})
        self.assertEqual(self.wb.sheetnames, list(NAMES))
        self.assertEqual(list(self.result), list(NAMES))

    def test_title_note_headers(self):
        for n in NAMES:
            ws = self.wb[n]
            self.assertEqual(ws["A2"].value, "T " + n)
            self.assertEqual(ws["A3"].value, "N " + n)
            self.assertEqual(tuple(c.value for c in ws[5]), HEADERS[n])

    def test_data_formula_and_formats(self):
        ws = self.wb[NAMES[1]]
        self.assertEqual(ws["A6"].value, "T001")
        self.assertEqual(ws["G6"].value, 100)
        self.assertEqual(ws["G6"].number_format, "#,##0")
        self.assertEqual(ws["I6"].value, "=" + NR)
        self.assertEqual(ws["I6"].number_format, "0.00%")

    def test_tables_freeze_width(self):
        refs = ("A5:AC6", "A5:AD6", "A5:AC6", "A5:K7")
        for i, n in enumerate(NAMES, 1):
            ws = self.wb[n]
            self.assertEqual(ws.tables["MetadataTable%d" % i].ref, refs[i - 1])
            self.assertEqual(ws.freeze_panes, "C6")
            self.assertEqual(ws.column_dimensions["A"].width, 18)

    def test_full_calc_on_load(self):
        with zipfile.ZipFile(self.path) as zf:
            self.assertRegex(zf.read("xl/workbook.xml").decode(), r'<calcPr[^>]*fullCalcOnLoad="1"')

    def test_refuses_existing_path(self):
        with self.assertRaises(ContractError) as ctx:
            write_workbook(self.path, sheets=make_sheets())
        self.assertEqual(str(ctx.exception), "output_exists")

