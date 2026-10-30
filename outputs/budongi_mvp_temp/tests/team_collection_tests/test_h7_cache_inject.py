"""Visible tests for H7 workbook/cache_inject.py (typical path)."""
from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl

from budongi.team_collection.contract import ContractError
from budongi.team_collection.workbook.cache_inject import inject_cached_values, read_cached


def _build(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = 1
    ws["A2"] = 2
    ws["B1"] = "=SUM(A1:A2)"
    ws["B2"] = "=A1*2.5"
    ws["C1"] = '=A1&"x"'
    ws["D1"] = "=A1>0"
    data = wb.create_sheet("Data")
    data["A1"] = "=1+1"
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


class CacheInjectVisibleTests(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        self.src = self.dir / "src.xlsx"
        self.dst = self.dir / "dst.xlsx"
        _build(self.src)

    def test_inject_returns_count_and_openpyxl_reads_cached_values(self) -> None:
        values = {("요약", "B1"): 3, ("요약", "B2"): 2.5, ("요약", "C1"): "1x",
                  ("요약", "D1"): True, ("Data", "A1"): 2}
        self.assertEqual(inject_cached_values(self.src, self.dst, values), 5)
        wb = openpyxl.load_workbook(self.dst, data_only=True)
        ws = wb["요약"]
        self.assertEqual(ws["B1"].value, 3)
        self.assertEqual(ws["B2"].value, 2.5)
        self.assertEqual(ws["C1"].value, "1x")
        self.assertIs(ws["D1"].value, True)
        self.assertEqual(wb["Data"]["A1"].value, 2)
        wb.close()
        wb = openpyxl.load_workbook(self.dst, data_only=False)
        self.assertEqual(wb["요약"]["B1"].value, "=SUM(A1:A2)")
        self.assertEqual(wb["Data"]["A1"].value, "=1+1")
        wb.close()

    def test_read_cached_reports_formula_value_and_type(self) -> None:
        before = read_cached(self.src)
        self.assertEqual(before[("요약", "B1")],
                         {"formula": "SUM(A1:A2)", "value": None, "type": None})
        self.assertEqual(set(before), {("요약", "B1"), ("요약", "B2"), ("요약", "C1"),
                                       ("요약", "D1"), ("Data", "A1")})
        inject_cached_values(self.src, self.dst, {("요약", "B1"): 3, ("요약", "C1"): "1x",
                                                  ("요약", "D1"): False})
        after = read_cached(self.dst)
        self.assertEqual(after[("요약", "B1")],
                         {"formula": "SUM(A1:A2)", "value": "3", "type": None})
        self.assertEqual(after[("요약", "C1")],
                         {"formula": 'A1&"x"', "value": "1x", "type": "str"})
        self.assertEqual(after[("요약", "D1")],
                         {"formula": "A1>0", "value": "0", "type": "b"})

    def test_untargeted_members_are_byte_identical_in_original_order(self) -> None:
        inject_cached_values(self.src, self.dst, {("요약", "B1"): 3})
        with zipfile.ZipFile(self.src) as zs, zipfile.ZipFile(self.dst) as zd:
            self.assertEqual(zd.namelist(), zs.namelist())
            self.assertIsNone(zd.testzip())
            for name in zs.namelist():
                if name != "xl/worksheets/sheet1.xml":
                    self.assertEqual(zd.read(name), zs.read(name), name)
            self.assertIn(b'fullCalcOnLoad="1"', zd.read("xl/workbook.xml"))

    def test_target_cell_without_formula_raises_and_writes_nothing(self) -> None:
        with self.assertRaises(ContractError) as cm:
            inject_cached_values(self.src, self.dst, {("요약", "A1"): 5})
        self.assertEqual(str(cm.exception), "cell_without_formula:요약!A1")
        self.assertFalse(self.dst.exists())


if __name__ == "__main__":
    unittest.main()
