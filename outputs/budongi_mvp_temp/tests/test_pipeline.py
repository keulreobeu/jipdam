"""Pipeline contracts: source evidence, exact values, isolated execution and recovery."""
from __future__ import annotations

import copy
import csv
import io
import hashlib
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from pathlib import Path

from openpyxl import Workbook

from budongi.pipeline.adapters import PRICE_HEADER, Output
from budongi.pipeline.benchmark import assess_targets, benchmark, compare, performance_report
from budongi.pipeline.catalog import Catalog, import_workbook, inventory_status, register_inventory
from budongi.pipeline.common import PipelineError, atomic_json, dumps, sha_file, signature, writer_lock
from budongi.pipeline.runner import lineage, plan_request, run_pipeline, status


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name)
        self.root = self.path / "pipeline"

    def bind(self, path, adapter, contract=None, descriptor=None, did="synthetic", iid="input", version="fixture_v1"):
        register_inventory(self.root, {"datasets": [{"dataset_id": did, "adapter_id": adapter,
            "mapping_version": version, "contract": contract or {}}], "files": [{"input_id": iid,
            "dataset_id": did, "mapping_version": version, "path": str(path), "sha256": sha_file(path),
            "descriptor": descriptor or {}}]})
        return {"dataset_id": did, "adapter_id": adapter, "mapping_version": version, "batch_id": "batch",
                "input_refs": [iid], "output_root": str(self.root)}

    def price(self, invalid=False):
        path = self.path / "price.zip"
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(PRICE_HEADER)
        for bjd, amount, code in [("1111010100", "9007199254740993", "0001"), ("2611010100", "0", "0002"), ("1111010100", "-1" if invalid else "0", "0001")]:
            row = {k: "" for k in PRICE_HEADER}
            row.update({"기준연도": "2026", "기준월": "01", "법정동코드": bjd, "공시가격": amount,
                        "전용면적": "84.1200", "단지코드": code, "동코드": "01", "호코드": "001", "단지명": "합성"})
            writer.writerow([row[k] for k in PRICE_HEADER])
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("price.csv", stream.getvalue().encode("utf-8"))
        return path

    def table(self, result):
        import pyarrow.parquet as pq
        return pq.read_table(self.root / "outputs" / result["run_id"] / "data.parquet").to_pylist()

    def test_catalog_revision_namespace_cells_and_unmeasured_zero(self):
        p = self.path / "metadata.xlsx"
        w = Workbook()
        s = w.active
        s.title = "데이터셋 목록"
        s["B2"] = "표 제목"
        s.append([])
        s["A5"], s["B5"], s["H5"] = "데이터셋_ID", "데이터셋명", "전체행수"
        s["A6"], s["B6"], s["H6"] = "M001", "합성", 0
        s["A7"], s["B7"] = "M002", "미측정"
        w.create_sheet("전처리 및 이슈 관리")["A2"] = "=1+1"
        w.save(p)
        a = import_workbook(self.root, p, "owner1")
        self.assertEqual(a["entries"], 2)
        self.assertTrue(import_workbook(self.root, p, "owner1")["reused"])
        self.assertFalse(import_workbook(self.root, p, "owner2")["reused"])
        s["B6"] = "개정"
        w.save(p)
        b = import_workbook(self.root, p, "owner1")
        self.assertEqual(b["changed"], ["M001"])
        with Catalog(self.root, readonly=True) as c:
            row = c.db.execute("SELECT cells_json FROM catalog_entries WHERE snapshot_id=? AND entry_id='M001'", (a["snapshot_id"],)).fetchone()
            self.assertEqual(json.loads(row[0])["H6"]["value"], 0)
            self.assertEqual(c.db.execute("SELECT count(*) FROM catalog_entries").fetchone()[0], 6)
        self.assertEqual(sha_file(self.root / "metadata" / sha_file(p) / "source.xlsx"), sha_file(p))

    def test_inventory_explicit_hash_missing_and_wrong_database(self):
        p = self.price()
        req = self.bind(p, "assessed_csv", descriptor={"member": "price.csv"})
        self.assertEqual(inventory_status(self.root)["inputs"][0]["status"], "verified")
        p.write_bytes(b"tamper")
        self.assertEqual(inventory_status(self.root)["inputs"][0]["status"], "hash_mismatch")
        with self.assertRaisesRegex(PipelineError, "input_missing_or_tampered"):
            plan_request(self.root, req)
        wrong = self.path / "wrong"
        wrong.mkdir()
        sqlite3.connect(wrong / "catalog.sqlite3").close()
        before = (wrong / "catalog.sqlite3").read_bytes()
        with self.assertRaisesRegex(PipelineError, "wrong_catalog"):
            Catalog(wrong)
        self.assertEqual(before, (wrong / "catalog.sqlite3").read_bytes())

    def test_assessed_exact_values_replay_batch_conflict_and_output_tamper(self):
        p = self.price()
        before = sha_file(p)
        req = self.bind(p, "assessed_csv", descriptor={"member": "price.csv"})
        result = run_pipeline(self.root, req)
        self.assertEqual(result["counts"], {"input_rows": 3, "accepted_input_rows": 2, "excluded_rows": 1,
                         "rejected_rows": 0, "output_rows": 2, "lineage_rows": 2})
        rows = self.table(result)
        self.assertEqual(rows[0]["amount_krw"], 9007199254740993)
        self.assertEqual(rows[0]["area_sqm"], "84.1200")
        self.assertIn("0001:01:001", rows[0]["source_subject_id"])
        self.assertEqual(rows[1]["amount_krw"], 0)
        self.assertEqual(lineage(self.root, result["run_id"], rows[0]["record_id"])[0]["source_row"], 1)
        self.assertTrue(run_pipeline(self.root, req)["reused"])
        with self.assertRaisesRegex(PipelineError, "batch_input_changed"):
            run_pipeline(self.root, req | {"max_input_rows": 1})
        self.assertEqual(sha_file(p), before)
        data = self.root / "outputs" / result["run_id"] / "data.parquet"
        data.write_bytes(b"tamper")
        with self.assertRaisesRegex(PipelineError, "output_missing_or_tampered"):
            run_pipeline(self.root, req)

    def test_failed_values_not_published_and_writer_is_exclusive(self):
        req = self.bind(self.price(invalid=True), "assessed_csv", descriptor={"member": "price.csv"})
        with writer_lock(self.root):
            with self.assertRaisesRegex(PipelineError, "writer_busy"):
                run_pipeline(self.root, req)
        with self.assertRaisesRegex(PipelineError, "required_values_rejected"):
            run_pipeline(self.root, req)
        self.assertEqual(status(self.root)[0]["status"], "failed")
        self.assertFalse((self.root / "outputs").exists())

    def test_statistics_one_input_row_many_cells_and_missing_zero(self):
        path = self.path / "아파트 매매평균가격.xlsx"
        w = Workbook()
        w.active.append(["지역명", "2026-08-01", "2026-09-01"])
        w.active.append(["서울", 0, "-"])
        w.save(path)
        req = self.bind(path, "statistics_xlsx")
        result = run_pipeline(self.root, req)
        self.assertEqual(result["counts"]["input_rows"], 1)
        self.assertEqual(result["counts"]["output_rows"], 2)
        rows = self.table(result)
        self.assertEqual([r["value_status"] for r in rows], ["numeric", "missing"])
        self.assertEqual([r["value_decimal"] for r in rows], ["0", None])
        self.assertEqual([r["source_column"] for r in rows], [2, 3])

    def test_molit_existing_normalizer_scope_and_source_namespace(self):
        path = self.path / "page.xml"
        items = []
        for day in [1, 20]:
            row = {"sggCd": "11110", "dealYear": "2026", "dealMonth": "10", "dealDay": str(day),
                   "aptNm": "합성", "deposit": "50,000", "monthlyRent": "0", "excluUseAr": "84.1"}
            items.append("<item>" + "".join(f"<{k}>{v}</{k}>" for k, v in row.items()) + "</item>")
        path.write_text("<response><header><resultCode>00</resultCode></header><body><pageNo>1</pageNo><numOfRows>1000</numOfRows><totalCount>2</totalCount><items>" + "".join(items) + "</items></body></response>", encoding="utf-8")
        req = self.bind(path, "molit_xml", {"observation_from": "2026-10-01", "snapshot_date": "2026-10-15"},
                        {"service": "rental", "page": 1, "per_page": 1000, "district": "11110", "month": "202610", "source_ref": "rental/page.xml"})
        result = run_pipeline(self.root, req)
        self.assertEqual(result["counts"]["accepted_input_rows"], 1)
        self.assertEqual(result["counts"]["excluded_rows"], 1)
        self.assertEqual(self.table(result)[0]["deposit_krw"], 500000000)
        self.assertEqual(self.table(result)[0]["source"], "molit_apt_rent")

    def test_neis_known_rejections_preserved_and_incomplete_pages_fail(self):
        path = self.path / "school.json"
        rows = [{"SD_SCHUL_CODE": "0001", "ATPT_OFCDC_SC_CODE": "B10", "SCHUL_NM": "학교", "SCHUL_KND_SC_NM": "초등학교", "FOND_YMD": "20260230"},
                {"ATPT_OFCDC_SC_CODE": "V10"}, {"ATPT_OFCDC_SC_CODE": "B10"},
                {"ATPT_OFCDC_SC_CODE": "B10", "SCHUL_KND_SC_NM": "초등학교", "SCHUL_NM": "코드없음"}]
        path.write_text(json.dumps({"schoolInfo": [{"head": [{"list_total_count": 4}, {"RESULT": {"CODE": "INFO-000"}}]}, {"row": rows}]}), encoding="utf-8")
        req = self.bind(path, "neis_json", {"expected_total_rows": 4}, {"page": 1, "expected_rows": 4})
        result = run_pipeline(self.root, req)
        self.assertEqual(result["counts"]["accepted_input_rows"], 1)
        self.assertEqual(result["counts"]["excluded_rows"], 2)
        self.assertEqual(result["counts"]["rejected_rows"], 1)
        row = self.table(result)[0]
        self.assertIsNone(row["founded_date"])
        self.assertIsNone(json.loads(row["payload_json"])["latitude"])
        quarantine = (self.root / "outputs" / result["run_id"] / "quarantine.jsonl").read_text(encoding="utf-8").splitlines()
        self.assertEqual(len(quarantine), 3)
        self.bind(path, "neis_json", {"expected_total_rows": 5}, {"page": 1, "expected_rows": 4}, did="incomplete", iid="incomplete")
        with self.assertRaisesRegex(PipelineError, "neis_total"):
            run_pipeline(self.root, req | {"dataset_id": "incomplete", "input_refs": ["incomplete"]})

    def test_csv_engine_comparisons_preserve_row_order_and_null_semantics(self):
        req = self.bind(self.price(), "assessed_csv", descriptor={"member": "price.csv"})
        results = [run_pipeline(self.root, req | {"batch_id": e, "engine": e}) for e in ["python", "pandas_full", "pandas_chunked", "duckdb"]]
        self.assertEqual(len({r["result_digest"] for r in results}), 1)
        self.assertEqual(len({r["semantic_signature"] for r in results}), 1)

    def test_crash_after_publish_recovers_receipt_and_before_publish_retries(self):
        req = self.bind(self.price(), "assessed_csv", descriptor={"member": "price.csv"})
        result = run_pipeline(self.root, req)
        with Catalog(self.root) as c, c.db:
            c.db.execute("UPDATE runs SET status='running',result_json=NULL")
        recovered = run_pipeline(self.root, req)
        self.assertTrue(recovered["recovered"])
        self.assertEqual(recovered["run_id"], result["run_id"])
        with Catalog(self.root) as c, c.db:
            c.db.execute("UPDATE runs SET status='interrupted'")
            old = dict(c.db.execute("SELECT * FROM runs").fetchone())
            old.update(run_id="deadbeef", status="running", result_json=None, receipt_sha256=None)
            c.db.execute("INSERT INTO runs VALUES(?,?,?,?,?,?,?,?,?,?,?)", tuple(old.values()))
        fresh = run_pipeline(self.root, req)
        self.assertEqual(fresh["attempt"], 3)
        self.assertEqual(fresh["result_digest"], result["result_digest"])
        self.assertIn("interrupted", {r["status"] for r in status(self.root)})

    def test_measurements_comparison_refuses_changed_results_and_no_baseline(self):
        self.assertEqual(performance_report(self.root)["baseline_status"], "baseline_unavailable")
        sample = {"duration_ms": 10, "memory": {"peak_rss_bytes": None}, "run_id": "a", "semantic_signature": "a", "result_digest": "a", "counts": {"output_rows": 1}, "environment": {"python": "same"}}
        baseline = [copy.deepcopy(sample) for _ in range(5)]
        candidate = copy.deepcopy(baseline)
        self.assertEqual(compare(baseline, candidate)["duration_improvement_percent"], 0)
        self.assertIsNone(compare(baseline, candidate)["baseline"]["max_peak_rss_bytes"])
        candidate[-1]["result_digest"] = "different"
        self.assertEqual(compare(baseline, candidate)["status"], "incomparable")
        self.assertNotIn("duration_improvement_percent", compare(baseline, candidate))

    def test_optimized_serialization_preserves_exact_event_ids_and_payload(self):
        directory = self.path / "serialization"
        directory.mkdir()
        dataset = {"dataset_id": "quote\"한글", "mapping_version": "v\\1", "adapter_id": "statistics_xlsx"}
        source = {"input_id": "i\n", "sha256": "0" * 64, "descriptor": {"source_ref": "ref\"\\"}}
        record = {"nested": {"한글": [None, True, "quote\"\\\n", 9007199254740993]}, "value_decimal": "0.000"}
        output = Output(directory, dataset)
        expected = hashlib.sha256()
        for row, status_, reason in [(1, "accepted", None), (2, "excluded", "이유\""), (3, "rejected", None)]:
            locator = {"input_id": source["input_id"], "source_sha256": source["sha256"], "source_ref": source["descriptor"]["source_ref"],
                       "section": "시트\"", "source_row": row, "source_column": 2}
            expected.update((dumps([locator, status_, reason, record]) + "\n").encode("utf-8"))
            output.emit(source, "시트\"", row, record, status=status_, reason=reason, column=2)
        output.close()
        import pyarrow.parquet as pq
        rows = pq.read_table(directory / "data.parquet").to_pylist()
        accepted_locator = dict(locator, source_row=1)
        self.assertEqual(rows[0]["record_id"], signature([dataset["dataset_id"], dataset["mapping_version"], accepted_locator]))
        self.assertEqual(rows[0]["payload_json"], dumps(record))
        self.assertEqual(output.digest.hexdigest(), expected.hexdigest())

    def test_target_assessment_requires_distinct_anchored_runs_and_exact_outputs(self):
        req = self.bind(self.price(), "assessed_csv", descriptor={"member": "price.csv"})
        before = benchmark(self.root, req)
        after = benchmark(self.root, req)
        target = {"dataset_id": "synthetic", "baseline_id": before["benchmark_id"], "candidate_id": after["benchmark_id"],
                  "source_signature": before["source_signature"], "baseline_sha256": sha_file(self.root / "benchmarks" / (before["benchmark_id"] + ".json")),
                  "max_duration_ratio": 100, "max_memory_ratio": 100}
        request = {"assessment_id": "approved-fixture", "approved_by": "synthetic", "approval_ref": "test only", "targets": [target]}
        result = assess_targets(self.root, request)
        self.assertEqual(result["targets_status"], "verified")
        self.assertTrue(result["outcomes"][0]["exact_parquet_and_quarantine"])
        self.assertIsNone(result["m1_complete"])
        self.assertEqual(performance_report(self.root)["targets_status"], "verified")
        self.assertEqual(assess_targets(self.root, request | {"assessment_id": "strict", "targets": [target | {"max_duration_ratio": 0.0001}]})["targets_status"], "not_met")
        self.assertEqual(performance_report(self.root)["targets_status"], "select_assessment_id")
        self.assertEqual(performance_report(self.root, "approved-fixture")["targets_status"], "verified")
        self.assertEqual(performance_report(self.root, "strict")["targets_status"], "not_met")
        with self.assertRaisesRegex(PipelineError, "approved_scope_mismatch"):
            assess_targets(self.root, request | {"targets": [target | {"source_signature": "f" * 64}]})
        benchmark_path = self.root / "benchmarks" / (before["benchmark_id"] + ".json")
        original = benchmark_path.read_bytes()
        benchmark_path.write_bytes(original + b"\n")
        with self.assertRaisesRegex(PipelineError, "approved_baseline_changed"):
            assess_targets(self.root, request)
        benchmark_path.write_bytes(original)
        with self.assertRaisesRegex(PipelineError, "distinct_baseline"):
            assess_targets(self.root, request | {"targets": [target | {"candidate_id": before["benchmark_id"]}]})
        rid = after["baseline"]["run_ids"][0]
        (self.root / "outputs" / rid / "data.parquet").write_bytes(b"tampered")
        with self.assertRaisesRegex(PipelineError, "output_missing_or_tampered"):
            performance_report(self.root)

    def test_cli_json_exit_codes_and_no_arbitrary_dispatch(self):
        req = self.bind(self.price(invalid=True), "assessed_csv", descriptor={"member": "price.csv"})
        path = self.path / "request.json"
        atomic_json(path, req)
        argv = [sys.executable, "-X", "utf8", "-m", "budongi.cli", "pipeline", "--root", str(self.root)]
        ready = subprocess.run(argv + ["plan", "--request", str(path)], capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(ready.returncode, 0, ready.stderr)
        failed = subprocess.run(argv + ["run", "--request", str(path)], capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(failed.returncode, 1)
        self.assertEqual(json.loads(failed.stdout)["status"], "failed")
        atomic_json(path, req | {"shell": "should not execute"})
        blocked = subprocess.run(argv + ["plan", "--request", str(path)], capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(blocked.returncode, 2)
        self.assertEqual(json.loads(blocked.stdout)["error_code"], "unknown_request_field")

    def test_malformed_region_code_blocks_publication_instead_of_exclusion(self):
        p = self.price()
        with zipfile.ZipFile(p) as z:
            payload = z.read("price.csv").decode("utf-8").replace("1111010100", "", 1)
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("price.csv", payload)
        req = self.bind(p, "assessed_csv", descriptor={"member": "price.csv"})
        with self.assertRaisesRegex(PipelineError, "required_values_rejected"):
            run_pipeline(self.root, req)

    def test_recovery_rejects_rewritten_receipt_after_coordinator_crash(self):
        req = self.bind(self.price(), "assessed_csv", descriptor={"member": "price.csv"})
        result = run_pipeline(self.root, req)
        directory = self.root / "outputs" / result["run_id"]
        data = directory / "data.parquet"
        data.write_bytes(data.read_bytes() + b"changed")
        rewritten = copy.deepcopy(result)
        rewritten["files"]["data.parquet"] = sha_file(data)
        atomic_json(directory / "run.json", rewritten)
        with Catalog(self.root) as c, c.db:
            c.db.execute("UPDATE runs SET status='running',result_json=NULL,receipt_sha256=NULL")
        with self.assertRaisesRegex(PipelineError, "receipt"):
            run_pipeline(self.root, req)

    def test_output_symlink_is_rejected_before_first_publication(self):
        req = self.bind(self.price(), "assessed_csv", descriptor={"member": "price.csv"})
        outside = self.path / "outside"
        outside.mkdir()
        # Windows test identity cannot create symlinks/junctions. Simulate only the OS query.
        with patch("budongi.pipeline.common.Path.is_symlink", autospec=True, side_effect=lambda p: p == self.root / "outputs"):
            with self.assertRaisesRegex(PipelineError, "symlink"):
                run_pipeline(self.root, req)
        self.assertEqual(list(outside.iterdir()), [])

    def test_cli_lock_contention_is_blocked_exit_two(self):
        req = self.bind(self.price(), "assessed_csv", descriptor={"member": "price.csv"})
        path = self.path / "request.json"
        atomic_json(path, req)
        with writer_lock(self.root):
            result = subprocess.run([sys.executable, "-m", "budongi.cli", "pipeline", "--root", str(self.root), "run", "--request", str(path)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(json.loads(result.stdout)["status"], "blocked")

    def test_retry_does_not_overlap_surviving_worker(self):
        import psutil
        req = self.bind(self.price(), "assessed_csv", descriptor={"member": "price.csv"})
        p = plan_request(self.root, req)
        work = self.root / "work" / "surviving"
        work.mkdir(parents=True)
        atomic_json(work / "job.json", {})
        atomic_json(work / "worker-state.json", {"pid": os.getpid(), "create_time": psutil.Process().create_time()})
        with Catalog(self.root) as c, c.db:
            c.db.execute("INSERT INTO runs VALUES(?,?,?,?,?,?,?,?,?,?,?)", ("surviving", "synthetic", "batch", "fixture_v1", p["input_signature"], 1, "running", "2026-10-29", None, None, None))
        with self.assertRaisesRegex(PipelineError, "worker_still_active"):
            run_pipeline(self.root, req)
        with self.assertRaisesRegex(PipelineError, "worker_still_active"):
            run_pipeline(self.root, req | {"batch_id": "another-batch"})

    def test_crash_before_worker_identity_cancels_delayed_launch_and_recovers(self):
        req = self.bind(self.price(), "assessed_csv", descriptor={"member": "price.csv"})
        p = plan_request(self.root, req)
        work = self.root / "work" / "delayed"
        artifacts = work / "artifacts"
        artifacts.mkdir(parents=True)
        atomic_json(work / "job.json", {"worker_protocol": 1, "directory": str(artifacts)})
        with Catalog(self.root) as c, c.db:
            c.db.execute("INSERT INTO runs VALUES(?,?,?,?,?,?,?,?,?,?,?)", ("delayed", "synthetic", "batch", "fixture_v1", p["input_signature"], 1, "running", "2026-10-29", None, None, None))
        result = run_pipeline(self.root, req | {"batch_id": "next"})
        self.assertEqual(result["status"], "success")
        self.assertTrue((work / "launch-cancelled.json").exists())
        delayed = subprocess.run([sys.executable, "-m", "budongi.pipeline.worker", "--job", str(work / "job.json")], capture_output=True)
        self.assertEqual(delayed.returncode, 0)
        self.assertEqual(list(artifacts.iterdir()), [])


if __name__ == "__main__":
    unittest.main()
