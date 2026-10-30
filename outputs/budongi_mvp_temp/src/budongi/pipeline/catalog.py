"""Versioned workbook evidence and explicit, hash-checked local input bindings."""
from __future__ import annotations

import json
import shutil
import sqlite3
from pathlib import Path

from .common import (PipelineError, atomic_json, checked_path, dumps, identifier,
                     now, sha_file, signature, writer_lock)

APPLICATION_ID = 0x4A504C31
SCHEMA = """
CREATE TABLE metadata_snapshots(snapshot_id TEXT PRIMARY KEY,namespace TEXT NOT NULL,sha256 TEXT NOT NULL,
 archive_ref TEXT NOT NULL,filename TEXT NOT NULL,created_at TEXT NOT NULL, UNIQUE(namespace,sha256));
CREATE TABLE catalog_entries(snapshot_id TEXT NOT NULL REFERENCES metadata_snapshots,entry_id TEXT NOT NULL,
 sheet TEXT NOT NULL,row_number INTEGER NOT NULL,cells_json TEXT NOT NULL,
 acquisition_status TEXT NOT NULL DEFAULT 'unknown',parse_status TEXT NOT NULL DEFAULT 'not_checked',
 quality_status TEXT NOT NULL DEFAULT 'unreviewed',relation_status TEXT NOT NULL DEFAULT 'unreviewed',
 license_status TEXT NOT NULL DEFAULT 'unreviewed',PRIMARY KEY(snapshot_id,entry_id));
CREATE TABLE metadata_rows(snapshot_id TEXT NOT NULL REFERENCES metadata_snapshots,sheet TEXT NOT NULL,
 row_number INTEGER NOT NULL,cells_json TEXT NOT NULL,PRIMARY KEY(snapshot_id,sheet,row_number));
CREATE TABLE datasets(dataset_id TEXT NOT NULL,mapping_version TEXT NOT NULL,adapter_id TEXT NOT NULL,
 contract_json TEXT NOT NULL,catalog_refs_json TEXT NOT NULL,PRIMARY KEY(dataset_id,mapping_version));
CREATE TABLE artifacts(sha256 TEXT PRIMARY KEY,byte_count INTEGER NOT NULL CHECK(byte_count>=0));
CREATE TABLE artifact_locations(location_id TEXT PRIMARY KEY,sha256 TEXT NOT NULL REFERENCES artifacts,
 path TEXT NOT NULL,checked_at TEXT NOT NULL,UNIQUE(sha256,path));
CREATE TABLE inputs(input_id TEXT PRIMARY KEY,dataset_id TEXT NOT NULL,mapping_version TEXT NOT NULL,
 sha256 TEXT NOT NULL REFERENCES artifacts,location_id TEXT NOT NULL REFERENCES artifact_locations,
 descriptor_json TEXT NOT NULL,FOREIGN KEY(dataset_id,mapping_version) REFERENCES datasets);
CREATE TABLE runs(run_id TEXT PRIMARY KEY,dataset_id TEXT NOT NULL,batch_id TEXT NOT NULL,
 mapping_version TEXT NOT NULL,input_signature TEXT NOT NULL,attempt INTEGER NOT NULL,
 status TEXT NOT NULL,started_at TEXT NOT NULL,finished_at TEXT,result_json TEXT,receipt_sha256 TEXT);
CREATE INDEX run_batch ON runs(dataset_id,batch_id,mapping_version,started_at);
"""


class Catalog:
    def __init__(self, root, *, readonly=False):
        self.root = checked_path(root)
        path = checked_path(self.root / "catalog.sqlite3")
        exists = path.exists()
        if readonly and not exists:
            raise PipelineError("catalog_missing")
        if not exists:
            self.root.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(path.as_uri() + ("?mode=ro" if readonly else "?mode=rwc"), uri=True)
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA foreign_keys=ON")
        if not exists:
            self.db.executescript(SCHEMA)
            self.db.execute(f"PRAGMA application_id={APPLICATION_ID}")
            self.db.execute("PRAGMA user_version=1")
            self.db.commit()
        elif (self.db.execute("PRAGMA application_id").fetchone()[0] != APPLICATION_ID
              or self.db.execute("PRAGMA user_version").fetchone()[0] != 1):
            self.db.close()
            raise PipelineError("wrong_catalog_schema")

    def close(self):
        self.db.close()

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.close()


def import_workbook(root, path, namespace):
    from openpyxl import load_workbook
    identifier(namespace)
    path = checked_path(path)
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise PipelineError("metadata_workbook_required")
    digest = sha_file(path)
    snapshot = signature([namespace, digest])
    with writer_lock(root), Catalog(root) as catalog:
        old = catalog.db.execute("SELECT * FROM metadata_snapshots WHERE snapshot_id=?", (snapshot,)).fetchone()
        if old:
            archive = checked_path(catalog.root / old["archive_ref"])
            if not archive.is_file() or sha_file(archive) != digest:
                raise PipelineError("metadata_archive_tampered")
            return {"snapshot_id": snapshot, "entries": catalog.db.execute(
                "SELECT count(*) FROM catalog_entries WHERE snapshot_id=?", (snapshot,)).fetchone()[0], "reused": True}
        archive = checked_path(catalog.root / "metadata" / digest / "source.xlsx")
        archive.parent.mkdir(parents=True, exist_ok=True)
        if archive.exists() and sha_file(archive) != digest:
            raise PipelineError("metadata_archive_tampered")
        if not archive.exists():
            shutil.copyfile(path, archive)
        if sha_file(archive) != digest:
            raise PipelineError("source_changed_during_copy")
        workbook = load_workbook(archive, read_only=True, data_only=False, keep_links=False)
        entries, metadata = [], []
        try:
            for sheet in workbook:
                # Some delivered OOXML dimensions are stale. Read the actual XML extent.
                sheet.reset_dimensions()
                header = None
                dataset_sheet = "데이터셋 목록" in sheet.title
                for row_number, row in enumerate(sheet.iter_rows(), 1):
                    cells = {}
                    for cell in row:
                        if cell.value is not None:
                            value = cell.value
                            if hasattr(value, "isoformat"):
                                value = value.isoformat()
                            cells[cell.coordinate] = {"value": value, "type": cell.data_type}
                    if not cells:
                        continue
                    metadata.append((snapshot, sheet.title, row_number, dumps(cells)))
                    if not dataset_sheet:
                        continue
                    first = row[0].value
                    if first == "데이터셋_ID":
                        header = True
                        continue
                    if header and isinstance(first, str) and first.strip():
                        if row[0].data_type == "f":
                            raise PipelineError("formula_dataset_id")
                        entries.append((snapshot, identifier(first.strip()), sheet.title, row[0].row, dumps(cells)))
        finally:
            workbook.close()
        if not entries or len({e[1] for e in entries}) != len(entries):
            raise PipelineError("missing_or_duplicate_catalog_entries")
        previous = catalog.db.execute("SELECT snapshot_id FROM metadata_snapshots WHERE namespace=? ORDER BY created_at DESC LIMIT 1", (namespace,)).fetchone()
        old_entries = {} if not previous else {r["entry_id"]: r["cells_json"] for r in catalog.db.execute(
            "SELECT entry_id,cells_json FROM catalog_entries WHERE snapshot_id=?", (previous[0],))}
        with catalog.db:
            catalog.db.execute("INSERT INTO metadata_snapshots VALUES(?,?,?,?,?,?)", (snapshot, namespace, digest,
                               archive.relative_to(catalog.root).as_posix(), path.name, now()))
            catalog.db.executemany("INSERT INTO catalog_entries(snapshot_id,entry_id,sheet,row_number,cells_json) VALUES(?,?,?,?,?)", entries)
            catalog.db.executemany("INSERT INTO metadata_rows VALUES(?,?,?,?)", metadata)
        current = {e[1]: e[4] for e in entries}
        return {"snapshot_id": snapshot, "entries": len(entries), "reused": False,
                "added": sorted(current.keys() - old_entries.keys()), "removed": sorted(old_entries.keys() - current.keys()),
                "changed": sorted(k for k in current.keys() & old_entries.keys() if current[k] != old_entries[k])}


def register_inventory(root, request):
    from .adapters import ADAPTERS
    if not isinstance(request, dict) or not isinstance(request.get("datasets"), list) or not isinstance(request.get("files"), list):
        raise PipelineError("invalid_inventory")
    with writer_lock(root), Catalog(root) as catalog:
        with catalog.db:
            for source in request["datasets"]:
                did, version = identifier(source["dataset_id"]), identifier(source["mapping_version"])
                adapter = source["adapter_id"]
                if adapter not in ADAPTERS:
                    raise PipelineError("unregistered_adapter")
                contract = source.get("contract", {})
                refs = source.get("catalog_refs", [])
                if not isinstance(contract, dict) or not isinstance(refs, list):
                    raise PipelineError("invalid_dataset_contract")
                for ref in refs:
                    if not catalog.db.execute("SELECT 1 FROM catalog_entries WHERE snapshot_id=? AND entry_id=?", (ref["snapshot_id"], ref["entry_id"])).fetchone():
                        raise PipelineError("unknown_catalog_entry")
                values = (did, version, adapter, dumps(contract), dumps(refs))
                old = catalog.db.execute("SELECT * FROM datasets WHERE dataset_id=? AND mapping_version=?", (did, version)).fetchone()
                if old and tuple(old) != values:
                    raise PipelineError("dataset_contract_changed_use_new_version")
                catalog.db.execute("INSERT OR IGNORE INTO datasets VALUES(?,?,?,?,?)", values)
            checked = []
            for source in request["files"]:
                iid = identifier(source["input_id"])
                did, version = identifier(source["dataset_id"]), identifier(source["mapping_version"])
                digest = source["sha256"]
                if not isinstance(digest, str) or len(digest) != 64 or any(c not in "0123456789abcdef" for c in digest):
                    raise PipelineError("explicit_sha256_required")
                path = checked_path(source["path"])
                if path.is_relative_to(catalog.root / "outputs"):
                    raise PipelineError("output_as_input_forbidden")
                descriptor = source.get("descriptor", {})
                if not isinstance(descriptor, dict):
                    raise PipelineError("invalid_input_descriptor")
                if not path.is_file():
                    checked.append({"input_id": iid, "status": "missing", "expected_sha256": digest})
                    continue
                if sha_file(path) != digest:
                    raise PipelineError("input_hash_mismatch")
                catalog.db.execute("INSERT OR IGNORE INTO artifacts VALUES(?,?)", (digest, path.stat().st_size))
                lid = signature([digest, str(path)])
                catalog.db.execute("INSERT OR IGNORE INTO artifact_locations VALUES(?,?,?,?)", (lid, digest, str(path), now()))
                values = (iid, did, version, digest, lid, dumps(descriptor))
                old = catalog.db.execute("SELECT * FROM inputs WHERE input_id=?", (iid,)).fetchone()
                if old and tuple(old) != values:
                    raise PipelineError("input_binding_changed_use_new_id")
                catalog.db.execute("INSERT OR IGNORE INTO inputs VALUES(?,?,?,?,?,?)", values)
                checked.append({"input_id": iid, "status": "verified", "sha256": digest})
        result = {"files": checked, "datasets": len(request["datasets"]), "checked_at": now()}
        atomic_json(catalog.root / "inventories" / (signature(result) + ".json"), result)
        return result


def inventory_status(root):
    with Catalog(root, readonly=True) as catalog:
        sources = [dict(r) for r in catalog.db.execute("SELECT i.input_id,i.sha256,l.path FROM inputs i JOIN artifact_locations l USING(location_id)")]
        for source in sources:
            path = checked_path(source["path"])
            source["status"] = "missing" if not path.is_file() else "verified" if sha_file(path) == source["sha256"] else "hash_mismatch"
        return {"catalog_entries": catalog.db.execute("SELECT count(*) FROM catalog_entries").fetchone()[0],
                "snapshots": catalog.db.execute("SELECT count(*) FROM metadata_snapshots").fetchone()[0], "inputs": sources}
