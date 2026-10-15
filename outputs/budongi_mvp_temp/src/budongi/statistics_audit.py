"""Audit preserved statistics against original XLSX cells and rounded exports."""
from __future__ import annotations

import hashlib
import io
import json
import sqlite3
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

from .statistics_storage import APPLICATION_ID, HEADERS, StatisticsError, _json, _period, _xml_values


def _component(label: str) -> str:
    label = label.replace("(%)", "").replace(" ", "")
    return "공급적절" if label == "적절" else label


def audit_statistics(db: Path, input_paths: list[Path] | None = None) -> dict:
    db = Path(db)
    connection = sqlite3.connect(db.resolve().as_uri() + "?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        if connection.execute("PRAGMA application_id").fetchone()[0] != APPLICATION_ID:
            raise StatisticsError("Not a statistics database")
        originals = {p.name: p for p in input_paths or []}
        files = []
        for source in connection.execute("SELECT * FROM source_files ORDER BY dataset_id,file_id").fetchall():
            payload = (db.parent / source["archive_path"]).read_bytes()
            if hashlib.sha256(payload).hexdigest() != source["sha256"]:
                raise StatisticsError("Raw archive differs from database SHA-256")
            filenames = [r[0] for r in connection.execute(
                "SELECT DISTINCT provided_filename FROM batch_files WHERE file_id=?", (source["file_id"],))]
            for filename in filenames:
                if filename in originals and hashlib.sha256(originals[filename].read_bytes()).hexdigest() != source["sha256"]:
                    raise StatisticsError(f"Provided file changed since import: {filename}")
            profile = json.loads(source["profile_json"])
            width = len(HEADERS[profile["layout"]])
            cells = {(r["sheet_name"],r["source_row"],r["source_column"]): r for r in connection.execute(
                "SELECT * FROM stat_observations WHERE file_id=?", (source["file_id"],))}
            xml = _xml_values(payload)
            book = load_workbook(io.BytesIO(payload), data_only=False)
            expected, numeric, missing = 0, 0, 0
            periods = set()
            try:
                for sheet in book:
                    for row in sheet.iter_rows(min_row=2):
                        for cell in row[width:]:
                            expected += 1
                            record = cells.pop((sheet.title,cell.row,cell.column), None)
                            if record is None:
                                raise StatisticsError("Missing database source cell")
                            header = sheet.cell(1,cell.column).value
                            period = _period(header,profile["frequency"])
                            periods.add(period)
                            if record["raw_value_json"] != _json(cell.value) or record["source_header_raw"] != str(header) or record["period_key"] != period:
                                raise StatisticsError("Source cell/header/period differs from database")
                            if cell.value is None or cell.value == "-":
                                missing += 1
                                if record["value_status"] != "missing" or record["value_decimal"] is not None:
                                    raise StatisticsError("Source missing cell changed meaning")
                            else:
                                numeric += 1
                                original = Decimal(xml[sheet.title][cell.coordinate])
                                if record["value_status"] != "numeric" or Decimal(record["value_decimal"]) != original:
                                    raise StatisticsError("Source numeric XML precision differs from database")
                if cells:
                    raise StatisticsError("Database has extra source cells")
            finally:
                book.close()
            files.append({"dataset_id":source["dataset_id"], "filenames":filenames,
                          "sha256":source["sha256"], "expected_cells":expected, "numeric_cells":numeric,
                          "missing_cells":missing, "first_period":min(periods), "last_period":max(periods),
                          "user_reported_update_period":profile.get("user_reported_update_period"),
                          "profile_general_base_period":profile.get("user_reported_general_base_period"),
                          "value_audit":"all_cells_equal"})
        duplicate_pairs = []
        for frequency in ("weekly", "monthly"):
            for kind in ("buyer_advantage", "jeonse_supply"):
                detailed = f"kb_{frequency}_{kind}"
                dashboard = detailed + "_seoul_dashboard"
                target = {}
                for r in connection.execute("""SELECT o.*,s.dimensions_json FROM latest_stat_observations o
                     JOIN series s USING(series_id) WHERE s.dataset_id=?""", (detailed,)):
                    dims = json.loads(r["dimensions_json"])
                    if dims.get("지역") == "서울":
                        target[r["period_key"],_component(str(dims["구분"]))] = r
                compared, exact, rounded, unmatched, largest = 0,0,0,0,Decimal(0)
                for r in connection.execute("""SELECT o.*,s.dimensions_json FROM latest_stat_observations o
                     JOIN series s USING(series_id) WHERE s.dataset_id=?""", (dashboard,)):
                    label = _component(str(json.loads(r["dimensions_json"])["조사항목"]))
                    t = target.get((r["period_key"],label))
                    if t is None or r["value_status"] != t["value_status"]:
                        unmatched += 1
                        continue
                    compared += 1
                    if r["value_status"] == "missing":
                        exact += 1
                        rounded += 1
                        continue
                    a,b = Decimal(r["value_decimal"]),Decimal(t["value_decimal"])
                    largest = max(largest,abs(a-b))
                    exact += a == b
                    rounded += a == b.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
                if compared or unmatched:
                    duplicate_pairs.append({"dashboard_dataset":dashboard, "detailed_dataset":detailed,
                                            "compared_cells":compared, "exact_equal_cells":exact,
                                            "one_decimal_equal_cells":rounded, "unmatched_cells":unmatched,
                                            "max_absolute_difference":str(largest),
                                            "rounded_duplicate":bool(compared and compared==rounded and not unmatched),
                                            "usage":"choose detailed dataset for calculations; preserve dashboard raw"})
        base_evidence = []
        for did in ("kb_monthly_apt_sale_index", "kb_monthly_apt_jeonse_index",
                    "kb_weekly_apt_sale_index", "kb_weekly_apt_jeonse_index"):
            candidate = "2026-01-12" if "weekly" in did else "2026-01"
            values = [Decimal(r[0]) for r in connection.execute("""SELECT value_decimal FROM latest_stat_observations o
                JOIN series s USING(series_id) WHERE dataset_id=? AND period_key=? AND value_status='numeric'""", (did,candidate))]
            if values:
                base_evidence.append({"dataset_id":did,"period":candidate,"numeric_regions":len(values),
                                      "regions_equal_100":sum(v==100 for v in values),
                                      "status":"observed values, not a substituted source definition"})
        report = {"files":files,"file_count":len(files), "total_cells":sum(f["expected_cells"] for f in files),
                  "numeric_cells":sum(f["numeric_cells"] for f in files),
                  "missing_cells":sum(f["missing_cells"] for f in files),
                  "duplicate_pairs":duplicate_pairs,"index_base_evidence":base_evidence,
                  "integrity_check":connection.execute("PRAGMA integrity_check").fetchone()[0],
                  "foreign_key_violations":len(connection.execute("PRAGMA foreign_key_check").fetchall())}
        if report["integrity_check"] != "ok" or report["foreign_key_violations"]:
            raise StatisticsError("Statistics integrity audit failed")
        return report
    finally:
        connection.close()


def record_audit_reviews(db: Path, report: dict) -> int:
    """Append evidence only. Original source cells and source profiles are unchanged."""
    source_ref = "statistics-audit-sha256:" + hashlib.sha256(_json(report).encode()).hexdigest()
    connection = sqlite3.connect(db)
    try:
        if connection.execute("PRAGMA application_id").fetchone()[0] != APPLICATION_ID:
            raise StatisticsError("Not a statistics database")
        connection.execute("PRAGMA foreign_keys=ON")
        added = 0
        with connection:
            reviews = [(r["dashboard_dataset"],r) for r in report["duplicate_pairs"]]
            reviews += [(r["dataset_id"],r) for r in report["index_base_evidence"]]
            for did, review in reviews:
                if not connection.execute("SELECT 1 FROM metadata_reviews WHERE dataset_id=? AND source_ref=?", (did,source_ref)).fetchone():
                    connection.execute("INSERT INTO metadata_reviews(dataset_id,recorded_at_utc,source_ref,metadata_json) VALUES(?,?,?,?)",
                                       (did,datetime.now(timezone.utc).isoformat(),source_ref,_json(review)))
                    added += 1
        return added
    finally:
        connection.close()
