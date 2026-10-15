"""Offline statistics XLSX store. Source cells and import vintages stay intact."""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
import sqlite3
import uuid
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

APPLICATION_ID = 0x4A535431
SCHEMA_VERSION = 1
REGISTRY = Path(__file__).with_name("statistics_sources.json")
PROVINCES = {"서울", "부산", "광주", "대구", "인천", "대전", "울산", "세종", "경기",
             "충북", "충남", "경북", "경남", "제주", "강원", "전북", "전남", "(구)광주광역시", "(구)전라남도"}
AGGREGATES = {"전국", "수도권", "6개광역시", "5개광역시", "기타지방", "지방"}
HEADERS = {"region": ["지역명"], "survey": ["지역", "구분"],
           "seoul_dashboard": ["조사항목"], "leading": ["아파트 지수 구분"],
           "pir": ["지역", "주택분위", "소득분위"], "selected_region": ["구분"]}


class StatisticsError(ValueError):
    """A source/contract failure; never a successful empty import."""


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _digest(value: bytes | object) -> str:
    return hashlib.sha256(value if isinstance(value, bytes) else _json(value).encode()).hexdigest()


def _registry() -> dict:
    value = json.loads(REGISTRY.read_text(encoding="utf-8"))
    if value.get("schema_version") != "statistics_sources_v1":
        raise StatisticsError("Unsupported statistics source registry")
    return value


def _profile(path: Path, registry: dict) -> dict:
    stem = re.sub(r"_\d{8}$", "", path.stem)
    matches = [p for p in registry["profiles"] if p["filename_stem"] == stem]
    if path.suffix.lower() != ".xlsx" or len(matches) != 1:
        raise StatisticsError(f"Unregistered statistics file: {path.name}")
    return matches[0]


def _period(value: object, frequency: str) -> str:
    if isinstance(value, (datetime, date)):
        parsed = value.date() if isinstance(value, datetime) else value
    elif isinstance(value, str) and re.fullmatch(r"'\d{2}\.\d{2}", value):
        parsed = date(2000 + int(value[1:3]), int(value[4:6]), 1)
    elif isinstance(value, str) and re.fullmatch(r"\d{4}\.\d{2}\.\d{2}", value):
        try:
            parsed = date(*map(int, value.split(".")))
        except ValueError as exc:
            raise StatisticsError(f"Invalid period header: {value!r}") from exc
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value).date()
        except ValueError as exc:
            raise StatisticsError(f"Invalid period header: {value!r}") from exc
    else:
        raise StatisticsError(f"Invalid period header: {value!r}")
    return parsed.isoformat() if frequency == "weekly" else parsed.strftime("%Y-%m")


def _xml_values(payload: bytes) -> dict[str, dict[str, str]]:
    """Keep numeric XML text, rather than recovering decimals from binary floats."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    import posixpath
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {r.attrib["Id"]: r.attrib["Target"] for r in rels}
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        result = {}
        for sheet in workbook.findall("m:sheets/m:sheet", ns):
            target = targets[sheet.attrib[f"{{{rel_ns}}}id"]]
            entry = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            xml = ET.fromstring(archive.read(entry))
            result[sheet.attrib["name"]] = {
                c.attrib["r"]: c.find("m:v", ns).text
                for c in xml.findall("m:sheetData/m:row/m:c", ns)
                if c.attrib.get("t", "n") == "n" and c.find("m:v", ns) is not None
            }
        return result


def _number(value: object, xml_text: str | None) -> tuple[str, str | None, float | None]:
    if value is None or value == "-":
        return "missing", None, None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or xml_text is None:
        raise StatisticsError(f"Unsupported value: {value!r}")
    try:
        exact = Decimal(xml_text)
        numeric = float(exact)
        if not exact.is_finite() or not math.isfinite(numeric):
            raise InvalidOperation
    except (InvalidOperation, ValueError, OverflowError) as exc:
        raise StatisticsError("Non-finite or invalid numeric cell") from exc
    return "numeric", str(exact), numeric


def _rows(payload: bytes, profile: dict):
    xml_values = _xml_values(payload)
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    try:
        seen_series = set()
        for sheet in workbook:
            headers = [c.value for c in sheet[1]]
            layout = profile["layout"]
            expected = HEADERS[layout]
            if layout == "region" and profile["provider"].startswith("kb_public"):
                expected = ["구분"]
            width = len(expected)
            if headers[:width] != expected or len(headers) <= width:
                raise StatisticsError(f"Unexpected layout: {sheet.title}")
            periods = [_period(v, profile["frequency"]) for v in headers[width:]]
            if len(set(periods)) != len(periods) or periods != sorted(periods):
                raise StatisticsError("Duplicate/unordered normalized periods")
            parent, group = None, None
            for row in sheet.iter_rows(min_row=2):
                labels = [c.value for c in row[:width]]
                if any(v is None for v in labels):
                    raise StatisticsError(f"Missing row label: {sheet.title}!{row[0].coordinate}")
                if any(c.data_type == "f" for c in row):
                    raise StatisticsError("Formula cells are not statistics source values")
                dimensions = dict(zip(expected, labels))
                label = str(labels[0])
                if layout == "region":
                    if label in PROVINCES:
                        parent, group = label, None
                    elif label in AGGREGATES:
                        parent, group = None, None
                    elif label in {"강북14개구", "강남11개구"}:
                        group = label
                    dimensions["원천상위지역"] = parent
                    dimensions["원천권역"] = group
                if layout == "seoul_dashboard":
                    dimensions["지역"] = "서울"
                elif layout == "selected_region":
                    dimensions["지역"] = profile.get("selected_region")
                    dimensions["지역검수"] = "pending" if not profile.get("selected_region") else "user_reviewed"
                if layout == "pir" and any(type(v) is not int or not 1 <= v <= 5 for v in labels[1:]):
                    raise StatisticsError("Invalid PIR quintile")
                unit = profile["default_unit"]
                if layout in {"survey", "seoul_dashboard"}:
                    unit = "index" if "지수" in str(labels[-1]) else "percent"
                sid = _digest([profile["dataset_id"], dimensions, unit])
                if sid in seen_series:
                    raise StatisticsError("Ambiguous duplicate series within a sheet")
                seen_series.add(sid)
                values = []
                for col, period in zip(row[width:], periods):
                    state, exact, numeric = _number(col.value, xml_values[sheet.title].get(col.coordinate))
                    values.append((period, str(headers[col.column - 1]), col.row, col.column,
                                   _json(col.value), state, exact, numeric))
                yield sheet.title, sid, dimensions, unit, values
    finally:
        workbook.close()


SCHEMA = """
CREATE TABLE datasets(dataset_id TEXT PRIMARY KEY, provider TEXT NOT NULL, source_view TEXT NOT NULL,
 frequency TEXT NOT NULL CHECK(frequency IN ('weekly','monthly')), layout TEXT NOT NULL);
CREATE TABLE import_batches(sequence INTEGER PRIMARY KEY, batch_id TEXT UNIQUE NOT NULL,
 fingerprint TEXT NOT NULL, imported_at_utc TEXT NOT NULL);
CREATE TABLE source_files(file_id TEXT PRIMARY KEY, dataset_id TEXT NOT NULL REFERENCES datasets,
 sha256 TEXT NOT NULL, size_bytes INTEGER NOT NULL CHECK(size_bytes>0), archive_path TEXT NOT NULL,
 mapping_version TEXT NOT NULL, profile_json TEXT NOT NULL);
CREATE TABLE batch_files(batch_id TEXT NOT NULL REFERENCES import_batches(batch_id),
 file_id TEXT NOT NULL REFERENCES source_files, provided_filename TEXT NOT NULL,
 PRIMARY KEY(batch_id,file_id));
CREATE TABLE series(series_id TEXT PRIMARY KEY, dataset_id TEXT NOT NULL REFERENCES datasets,
 dimensions_json TEXT NOT NULL, unit_code TEXT NOT NULL, unit_status TEXT NOT NULL);
CREATE TABLE stat_observations(observation_id TEXT PRIMARY KEY, file_id TEXT NOT NULL REFERENCES source_files,
 series_id TEXT NOT NULL REFERENCES series, period_key TEXT NOT NULL, source_header_raw TEXT NOT NULL,
 sheet_name TEXT NOT NULL, source_row INTEGER NOT NULL CHECK(source_row>1),
 source_column INTEGER NOT NULL CHECK(source_column>1), raw_value_json TEXT NOT NULL,
 value_status TEXT NOT NULL CHECK(value_status IN ('numeric','missing')), value_decimal TEXT,
 value_numeric REAL, UNIQUE(file_id,sheet_name,source_row,source_column),
 CHECK((value_status='missing' AND value_decimal IS NULL AND value_numeric IS NULL) OR
       (value_status='numeric' AND value_decimal IS NOT NULL AND value_numeric IS NOT NULL)));
CREATE TABLE metadata_reviews(review_id INTEGER PRIMARY KEY, dataset_id TEXT NOT NULL REFERENCES datasets,
 recorded_at_utc TEXT NOT NULL, source_ref TEXT NOT NULL, metadata_json TEXT NOT NULL);
CREATE INDEX series_period ON stat_observations(series_id,period_key);
CREATE INDEX source_observations ON stat_observations(file_id);
CREATE VIEW latest_stat_observations AS
 SELECT * FROM (SELECT o.*, b.batch_id, b.imported_at_utc,
 ROW_NUMBER() OVER(PARTITION BY o.series_id,o.period_key ORDER BY b.sequence DESC,o.observation_id) AS version_rank
 FROM stat_observations o JOIN batch_files bf ON bf.file_id=o.file_id
 JOIN import_batches b ON b.batch_id=bf.batch_id) WHERE version_rank=1;
"""


def _connect(path: Path, *, new: bool = False) -> sqlite3.Connection:
    connection = sqlite3.connect(path)
    connection.row_factory = sqlite3.Row
    if new:
        connection.executescript(SCHEMA)
        connection.execute(f"PRAGMA application_id={APPLICATION_ID}")
        connection.execute(f"PRAGMA user_version={SCHEMA_VERSION}")
    elif (connection.execute("PRAGMA application_id").fetchone()[0] != APPLICATION_ID or
          connection.execute("PRAGMA user_version").fetchone()[0] != SCHEMA_VERSION):
        connection.close()
        raise StatisticsError("Target is not a statistics_v1 database")
    connection.execute("PRAGMA foreign_keys=ON")
    return connection


def _archive(root: Path, payload: bytes, sha: str) -> str:
    relative = Path("raw") / f"{sha}.xlsx"
    target = root / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        if _digest(target.read_bytes()) != sha:
            raise StatisticsError("Preserved source archive hash mismatch")
    else:
        try:
            with target.open("xb") as stream:
                stream.write(payload)
        except FileExistsError:
            if _digest(target.read_bytes()) != sha:
                raise StatisticsError("Concurrent source archive mismatch") from None
    return relative.as_posix()


def _summary(connection: sqlite3.Connection, batch_id: str | None = None) -> dict:
    where, params = ("WHERE bf.batch_id=?", [batch_id]) if batch_id else ("", [])
    files = [dict(r) for r in connection.execute(f"""
     SELECT bf.batch_id, bf.provided_filename, f.file_id, f.dataset_id, f.sha256, f.archive_path,
     count(o.observation_id) AS cells, sum(o.value_status='numeric') AS numeric_cells,
     sum(o.value_status='missing') AS missing_cells, min(o.period_key) AS first_period,
     max(o.period_key) AS last_period
     FROM batch_files bf JOIN source_files f USING(file_id) JOIN stat_observations o USING(file_id)
     {where} GROUP BY bf.batch_id,bf.file_id ORDER BY bf.provided_filename
    """, params)]
    return {"schema_version": "statistics_v1", "batch_id": batch_id, "files": files,
            "total_source_files": connection.execute("SELECT count(*) FROM source_files").fetchone()[0],
            "total_observations": connection.execute("SELECT count(*) FROM stat_observations").fetchone()[0],
            "total_latest_observations": connection.execute("SELECT count(*) FROM latest_stat_observations").fetchone()[0],
            "metadata_status": "user_supplied_pending_review"}


def _write_manifest(root: Path, connection: sqlite3.Connection, batch_id: str, fingerprint: str) -> None:
    directory = root / "manifests"
    directory.mkdir(exist_ok=True)
    target = directory / f"{batch_id}.json"
    files = _summary(connection, batch_id)["files"]
    manifest = {"schema_version": "statistics_v1", "batch_id": batch_id,
                "fingerprint": fingerprint, "files": files,
                "imported_at_utc": connection.execute(
                    "SELECT imported_at_utc FROM import_batches WHERE batch_id=?", (batch_id,)).fetchone()[0],
                "batch_cells": sum(f["cells"] for f in files), "metadata_status": "pending_review"}
    if target.exists():
        stored = json.loads(target.read_text(encoding="utf-8"))
        if stored != manifest:
            raise StatisticsError("Existing batch manifest differs from preserved batch")
        return
    temporary = directory / f".{batch_id}.{uuid.uuid4().hex}.tmp"
    try:
        temporary.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        try:
            _publish_new(temporary, target)
        except FileExistsError:
            if json.loads(target.read_text(encoding="utf-8")) != manifest:
                raise StatisticsError("Concurrent batch manifest mismatch") from None
    finally:
        temporary.unlink(missing_ok=True)


def _publish_new(temporary: Path, target: Path) -> None:
    """Windows rename refuses replacement; POSIX uses atomic exclusive linking."""
    import os
    if os.name == "nt":
        os.rename(temporary, target)
    else:
        os.link(temporary, target)


def import_statistics(db: Path, batch_id: str, paths: list[Path]) -> dict:
    db = Path(db)
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,100}", batch_id) or not paths:
        raise StatisticsError("An explicit safe batch ID and nonempty file list are required")
    registry = _registry()
    inputs, seen = [], set()
    for path in map(Path, paths):
        profile = _profile(path, registry)
        if profile["dataset_id"] in seen:
            raise StatisticsError("Only one vintage of each dataset may enter a batch")
        seen.add(profile["dataset_id"])
        payload = path.read_bytes()
        sha = _digest(payload)
        inputs.append((path, profile, payload, sha))
    fingerprint = _digest([registry["mapping_version"], sorted(
        (p.name, sha, _json(profile)) for p, profile, _, sha in inputs)])
    new = not db.exists()
    db.parent.mkdir(parents=True, exist_ok=True)
    working = db.with_name(f".{db.name}.{uuid.uuid4().hex}.tmp") if new else db
    connection = None
    try:
        connection = _connect(working, new=new)
        existing = connection.execute("SELECT fingerprint FROM import_batches WHERE batch_id=?", (batch_id,)).fetchone()
        if existing:
            if existing[0] != fingerprint:
                raise StatisticsError("Batch ID already belongs to different inputs/mapping")
            for _, _, payload, sha in inputs:
                _archive(db.parent, payload, sha)
            _write_manifest(db.parent, connection, batch_id, fingerprint)
            result = _summary(connection, batch_id)
            result["replayed"] = True
            result["new_observations"] = 0
            return result
        new_observations = 0
        with connection:
            connection.execute("INSERT INTO import_batches(batch_id,fingerprint,imported_at_utc) VALUES(?,?,?)",
                               (batch_id, fingerprint, datetime.now(timezone.utc).isoformat()))
            for path, profile, payload, sha in inputs:
                archived = _archive(db.parent, payload, sha)
                did = profile["dataset_id"]
                connection.execute("INSERT OR IGNORE INTO datasets VALUES(?,?,?,?,?)",
                                   (did, profile["provider"], profile["source_view"], profile["frequency"], profile["layout"]))
                fid = _digest([sha, did, registry["mapping_version"], profile])
                found = connection.execute("SELECT 1 FROM source_files WHERE file_id=?", (fid,)).fetchone()
                if not found:
                    connection.execute("INSERT INTO source_files VALUES(?,?,?,?,?,?,?)",
                                       (fid, did, sha, len(payload), archived, registry["mapping_version"], _json(profile)))
                    row_count = 0
                    for sheet, sid, dimensions, unit, values in _rows(payload, profile):
                        row_count += 1
                        connection.execute("INSERT OR IGNORE INTO series VALUES(?,?,?,?,?)",
                                           (sid, did, _json(dimensions), unit, profile["unit_status"]))
                        records = [( _digest([fid,sheet,r,c]), fid, sid, period, header, sheet, r, c,
                                    raw, state, exact, numeric)
                                   for period,header,r,c,raw,state,exact,numeric in values]
                        connection.executemany("INSERT INTO stat_observations VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", records)
                        new_observations += len(records)
                    if not row_count:
                        raise StatisticsError("Empty statistics workbook is not successful collection")
                connection.execute("INSERT INTO batch_files VALUES(?,?,?)", (batch_id, fid, path.name))
            if connection.execute("PRAGMA foreign_key_check").fetchall():
                raise StatisticsError("Statistics foreign key failure")
        result = _summary(connection, batch_id)
        result.update(replayed=False, new_observations=new_observations)
        if new:
            connection.close()
            connection = None
            # No concurrent importer may replace a database created since preflight.
            try:
                _publish_new(working, db)
            except FileExistsError:
                raise StatisticsError("Statistics database appeared during import; refusing replacement") from None
            connection = _connect(db)
        _write_manifest(db.parent, connection, batch_id, fingerprint)
        return result
    except (KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise StatisticsError("Invalid XLSX archive structure") from exc
    finally:
        if connection is not None:
            connection.close()
        if new and working.exists():
            working.unlink()


def statistics_status(db: Path) -> dict:
    connection = sqlite3.connect(Path(db).resolve().as_uri() + "?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        if connection.execute("PRAGMA application_id").fetchone()[0] != APPLICATION_ID:
            raise StatisticsError("Target is not a statistics_v1 database")
        result = _summary(connection)
        result["integrity_check"] = connection.execute("PRAGMA integrity_check").fetchone()[0]
        result["foreign_key_violations"] = len(connection.execute("PRAGMA foreign_key_check").fetchall())
        result["metadata_reviews"] = [dict(r) for r in connection.execute(
            "SELECT dataset_id,source_ref,metadata_json FROM metadata_reviews ORDER BY review_id")]
        return result
    finally:
        connection.close()
