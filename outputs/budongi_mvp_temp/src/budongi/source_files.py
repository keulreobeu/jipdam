"""Registry-approved CSV/JSON/XLSX raw preservation and relational staging."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import tempfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .collection import CollectionError, digest_file, read_json, safe_file, utc_now, write_json
from .rental_recommendation import SEOUL_DISTRICTS
from .rental_storage import observation_id_for, _assert_no_credentials

FILE_STAGE_VERSION = "registered_file_stage_v4"


def read_registry(path: Path | None = None) -> dict:
    registry = read_json(path or Path(__file__).with_name("source_registry.json"))
    sources = registry.get("sources")
    if registry.get("schema_version") != "source_registry_v1" or not isinstance(sources, list):
        raise CollectionError("invalid_source_registry")
    identifiers = [s.get("source_id") for s in sources if isinstance(s, dict)]
    if len(identifiers) != len(sources) or any(not isinstance(i, str) or not i.strip() for i in identifiers) or len(set(identifiers)) != len(identifiers):
        raise CollectionError("invalid_source_registry_ids")
    return registry


def _cell(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float) and math.isfinite(value):
        return value
    if isinstance(value, dict) and all(isinstance(k, str) for k in value):
        return {k: _cell(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_cell(v) for v in value]
    raise CollectionError("unsupported_file_cell")


def _content_key(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
                                    allow_nan=False).encode("utf-8")).hexdigest()


def _json_section(source):
    return "json:" + "".join("/" + str(v).replace("~", "~0").replace("/", "~1")
                            for v in source.get("items_path", []))


def _sections(path, source):
    if source["format"] != "xlsx":
        return [_json_section(source) if source["format"] == "json" else "table"]
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        selected = source.get("sheets", book.sheetnames)
        if (not isinstance(selected, list) or any(not isinstance(n, str) for n in selected)
                or len(selected) != len(set(selected)) or set(selected) != set(book.sheetnames)):
            raise CollectionError("all_workbook_sheets_must_be_declared")
        return selected
    finally:
        book.close()


def _table_rows(rows):
    iterator = iter(rows)
    headers = next(iterator, None)
    if headers is None:
        return
    headers = [str(v).strip() if v is not None else "" for v in headers]
    if not all(headers) or len(set(headers)) != len(headers):
        raise CollectionError("invalid_or_duplicate_headers")
    for position, values in enumerate(iterator, 2):
        values = list(values)
        if not any(v is not None and v != "" for v in values):
            continue
        if len(values) > len(headers):
            raise CollectionError("row_has_extra_columns")
        values += [None] * (len(headers) - len(values))
        yield position, {h: _cell(v) for h, v in zip(headers, values)}


def file_rows(path: Path, source: dict):
    format_name = source.get("format")
    if format_name == "csv":
        with path.open("r", encoding=source.get("encoding", "utf-8-sig"), newline="") as stream:
            for position, row in _table_rows(csv.reader(stream, delimiter=source.get("delimiter", ","))):
                yield "table", position, row
    elif format_name == "json":
        with path.open("r", encoding="utf-8-sig") as stream:
            # Floating-point JSON literals must not lose won precision before Decimal mapping.
            # Their lexical numeric text is retained; original bytes remain authoritative.
            value = json.load(stream, parse_float=str)
        for component in source.get("items_path", []):
            value = value[component]
        if not isinstance(value, list) or any(not isinstance(row, dict) for row in value):
            raise CollectionError("json_items_must_be_object_array")
        for index, row in enumerate(value, 1):
            yield _json_section(source), index, {k: _cell(v) for k, v in row.items()}
    elif format_name == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            requested = source.get("sheets", book.sheetnames)
            if not isinstance(requested, list) or len(requested) != len(set(requested)) or set(requested) != set(book.sheetnames):
                raise CollectionError("all_workbook_sheets_must_be_declared")
            for name in requested:
                for position, row in _table_rows(book[name].iter_rows(values_only=True)):
                    yield name, position, row
        finally:
            book.close()
    else:
        raise CollectionError("unsupported_registered_file_format")


def _map_fields(raw: dict, rules: dict) -> dict:
    mapped = {}
    for target, rule in rules.items():
        column = rule["column"]
        value = raw
        for component in ([column] if isinstance(column, str) else column):
            if not isinstance(value, dict):
                raise CollectionError("file_field_path_not_object")
            value = value.get(component)
            if value is None:
                break
        if value is None or isinstance(value, str) and not value.strip():
            if rule.get("required"):
                raise CollectionError("required_file_field_missing")
            mapped[target] = None
            continue
        kind = rule.get("type", "string")
        try:
            if kind == "string": result = str(value).strip()
            elif kind == "date":
                text = str(value).strip()
                result = (datetime.strptime(text, rule["format"]).date().isoformat() if rule.get("format")
                          else date.fromisoformat(text).isoformat() if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text)
                          else None)
                if result is None:
                    raise CollectionError("invalid_file_date")
            elif kind in {"integer", "number", "amount_krw"}:
                if kind == "amount_krw" and isinstance(value, float) and abs(value) >= 2**53:
                    raise CollectionError("unsafe_floating_currency_use_exact_text")
                text = str(value).strip()
                if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?(?:[eE][+-]?\d+)?", text):
                    raise CollectionError("invalid_file_number")
                number = Decimal(text.replace(",", ""))
                if kind == "amount_krw":
                    if rule.get("unit") not in {"krw", "manwon"}: raise CollectionError("unknown_amount_unit")
                    number *= 10000 if rule["unit"] == "manwon" else 1
                if not number.is_finite(): raise InvalidOperation
                if kind in {"integer", "amount_krw"}:
                    if number != number.to_integral_value() or not -(2**63) <= number < 2**63: raise InvalidOperation
                    result = int(number)
                else:
                    result = float(number)
                    if not math.isfinite(result): raise InvalidOperation
                if "min" in rule and result < rule["min"] or "max" in rule and result > rule["max"]:
                    raise CollectionError("file_field_out_of_range")
            else: raise CollectionError("unsupported_field_mapping_type")
        except CollectionError:
            raise
        except (ValueError, TypeError, InvalidOperation, OverflowError):
            reason = "invalid_file_date" if kind == "date" else "invalid_file_number" if kind in {"integer", "number", "amount_krw"} else "invalid_file_field"
            raise CollectionError(reason) from None
        mapped[target] = result
    return mapped


def normalize_file_record(row, source):
    """The registered record contract maps typed facts; never guess units or IDs."""
    required = {"sgg_cd": "string", "contract_date": "date", "apartment_name": "string", "area_sqm": "number"}
    rental = source.get("record_type") == "rental"
    if not rental and source.get("record_type") != "sale":
        raise CollectionError("unsupported_file_record_type")
    for field in (("deposit_krw", "monthly_rent_krw") if rental else ("price_krw",)):
        required[field] = "amount_krw"
    for field, kind in required.items():
        rule = source["field_mapping"].get(field, {})
        if rule.get("type", "string") != kind or not rule.get("required") or row.get(field) is None:
            raise CollectionError("file_target_requires_typed_field_" + field)
    if row["area_sqm"] <= 0 or any(row[field] < 0 for field in required if required[field] == "amount_krw"):
        raise CollectionError("file_target_negative_amount_or_nonpositive_area")
    if (rental and row.get("price_krw") is not None
            or not rental and (row.get("deposit_krw") is not None or row.get("monthly_rent_krw") is not None)):
        raise CollectionError("file_target_kind_amount_conflict")
    if row["sgg_cd"] not in source.get("regions", {code: name for name, code in SEOUL_DISTRICTS.items()}):
        raise CollectionError("file_target_region_outside_registered_scope")
    for field in ("floor", "build_year"):
        if row.get(field) is not None and source["field_mapping"][field].get("type") != "integer":
            raise CollectionError("file_target_requires_integer_" + field)
    for field in ("source_effective_from", "source_effective_to"):
        if row.get(field) is not None and source["field_mapping"][field].get("type") != "date":
            raise CollectionError("file_target_requires_date_" + field)
    if row.get("source_effective_from") and row.get("source_effective_to") and row["source_effective_from"] > row["source_effective_to"]:
        raise CollectionError("file_target_invalid_effective_period")
    kind = "jeonse" if rental and row["monthly_rent_krw"] == 0 else "wolse" if rental else "sale"
    if row.get("kind") is not None and row["kind"] != kind:
        raise CollectionError("file_target_kind_amount_conflict")
    row = dict(row)
    row["kind"] = kind
    for entity in ("complex", "building"):
        field = "source_" + entity + "_id"
        if row.get(field) is not None:
            contracts = source.get("official_identifiers", {})
            contract = contracts.get(entity, {}) if isinstance(contracts, dict) else {}
            if (not isinstance(contract, dict) or not isinstance(contract.get("namespace"), str)
                    or not contract["namespace"].strip() or not isinstance(contract.get("basis"), str)
                    or not contract["basis"].strip() or source["field_mapping"][field].get("type", "string") != "string"):
                raise CollectionError("file_identifier_requires_explicit_official_contract")
    address_fields = ("legal_dong_code", "main_number", "sub_number", "land_code")
    if any(row.get(field) is not None for field in address_fields):
        if not isinstance(source.get("address_basis"), str) or not source["address_basis"].strip():
            raise CollectionError("file_address_requires_official_component_basis")
        for field, length in zip(address_fields, (5, 4, 4, 1)):
            value = row.get(field)
            if (not isinstance(value, str) or not value.isascii() or not value.isdigit()
                    or not 1 <= len(value) <= length or field in {"legal_dong_code", "land_code"} and len(value) != length):
                raise CollectionError("file_address_invalid_component")
            row[field] = value.zfill(length)
    return row


def stage_file(*, registry_path: Path, source_id: str, input_path: Path, output: Path) -> dict:
    registry = read_registry(registry_path)
    source = next((s for s in registry["sources"] if s["source_id"] == source_id), None)
    if not source or source.get("status") != "approved" or not source.get("scope") or not source.get("approval_ref"):
        raise CollectionError("file_source_requires_approved_registry_scope")
    if source.get("format") not in {"csv", "json", "xlsx"} or not isinstance(source.get("field_mapping"), dict):
        raise CollectionError("file_source_requires_format_and_mapping")
    if not source.get("mapping_version") or not source["field_mapping"]:
        raise CollectionError("file_source_requires_mapping_version")
    expected = source.get("expected_rows")
    section_expected = source.get("expected_rows_by_section")
    if expected is not None and (type(expected) is not int or expected < 0):
        raise CollectionError("invalid_expected_file_rows")
    if section_expected is not None and (not isinstance(section_expected, dict)
            or any(not isinstance(k, str) or type(v) is not int or v < 0 for k, v in section_expected.items())):
        raise CollectionError("invalid_expected_section_rows")
    for target, rule in source["field_mapping"].items():
        if not isinstance(target, str) or not isinstance(rule, dict):
            raise CollectionError("invalid_file_field_mapping")
        column = rule.get("column")
        if not (isinstance(column, str) and column or isinstance(column, list) and column
                and all(isinstance(k, str) and k for k in column)):
            raise CollectionError("invalid_file_field_mapping")
    if not input_path.is_file() or input_path.is_symlink():
        raise CollectionError("invalid_source_file")
    digest = digest_file(input_path)
    mapping_hash = digest_file(registry_path)
    contract_hash = _content_key(source)
    if output.exists():
        old = read_json(output / "manifest.json")
        if old.get("schema_version") != FILE_STAGE_VERSION:
            raise CollectionError("file_stage_schema_change_use_new_directory")
        if old.get("state") == "parsing_failed":
            raise CollectionError("file_stage_parsing_failed_use_new_directory")
        if old.get("source_id") == source_id and old.get("raw_sha256") == digest and old.get("source_contract_sha256") == contract_hash:
            if digest_file(safe_file(output, old["raw_file"])) != digest or digest_file(output / "staging.sqlite3") != old["staging_sha256"]:
                raise CollectionError("file_stage_integrity_failed")
            return old | {"reused": True}
        raise CollectionError("file_stage_exists_use_new_directory")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix=".file-stage-", dir=output.parent))
    c = sqlite3.connect(temporary / "staging.sqlite3")
    c.executescript("""PRAGMA foreign_keys=ON;
      CREATE TABLE sources(source_id TEXT PRIMARY KEY,source_format TEXT NOT NULL,mapping_version TEXT NOT NULL);
      CREATE TABLE raw_files(file_id TEXT PRIMARY KEY,source_id TEXT NOT NULL REFERENCES sources,file_ref TEXT NOT NULL,
       sha256 TEXT NOT NULL,byte_count INTEGER NOT NULL CHECK(byte_count>=0),retrieved_at_utc TEXT NOT NULL);
      CREATE TABLE staging_items(observation_id TEXT PRIMARY KEY,file_id TEXT NOT NULL REFERENCES raw_files,
       section TEXT NOT NULL,row_position INTEGER NOT NULL CHECK(row_position>0),raw_json TEXT NOT NULL,
       standard_json TEXT,validation_status TEXT NOT NULL CHECK(validation_status IN('valid','rejected')),reason TEXT,
       UNIQUE(file_id,section,row_position));""")
    counts = {"valid": 0, "rejected": 0}
    sections = {}
    preserved = False
    try:
        # JSON selectors describe distinct logical file views, while mapping versions
        # reuse the same original row identities for an unchanged selector.
        prefix = "raw/" + (_content_key(source.get("items_path", [])) + "/" if source["format"] == "json" else "")
        raw_name = prefix + "source." + source["format"]
        (temporary / prefix).mkdir(parents=True)
        shutil.copyfile(input_path, temporary / raw_name)
        if digest_file(temporary / raw_name) != digest:
            raise CollectionError("source_changed_during_copy")
        preserved = True
        sections = {section: 0 for section in _sections(temporary / raw_name, source)}
        # Parse the preserved copy, not a changing upstream/local source.
        with c:
            c.execute("INSERT INTO sources VALUES (?,?,?)", (source_id, source["format"], source.get("mapping_version", "file_mapping_v1")))
            c.execute("INSERT INTO raw_files VALUES (?,?,?,?,?,?)", (digest, source_id, raw_name, digest,
                      (temporary / raw_name).stat().st_size, utc_now()))
            for section, position, raw in file_rows(temporary / raw_name, source):
                try:
                    _assert_no_credentials(raw)
                except ValueError:
                    raise CollectionError("credential_fields_forbidden_in_source_file") from None
                sections[section] = sections.get(section, 0) + 1
                state, reason, standard = "valid", None, None
                try:
                    standard = _map_fields(raw, source["field_mapping"])
                    if source.get("record_type"):
                        standard = normalize_file_record(standard, source)
                except CollectionError as exc: state, reason = "rejected", str(exc)
                counts[state] += 1
                key = observation_id_for(source=source_id, source_file_ref=_content_key(["section", section]) + "/" + raw_name,
                                         source_content_sha256=digest, source_row_number=position)
                c.execute("INSERT INTO staging_items VALUES (?,?,?,?,?,?,?,?)", (key,digest,section,position,
                          json.dumps(raw,ensure_ascii=False,sort_keys=True,allow_nan=False),
                          json.dumps(standard,ensure_ascii=False,sort_keys=True,allow_nan=False) if standard is not None else None,
                          state,reason))
        if c.execute("PRAGMA foreign_key_check").fetchall() or c.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise CollectionError("file_stage_integrity_failed")
        for table in ("sources", "raw_files", "staging_items"):
            for action in ("INSERT", "UPDATE", "DELETE"):
                c.execute(f"CREATE TRIGGER seal_{table}_{action} BEFORE {action} ON {table} BEGIN SELECT RAISE(ABORT,'immutable file staging'); END")
        c.commit()
        c.close(); c = None
        received = sum(counts.values())
        expected = source.get("expected_rows")
        section_expected = source.get("expected_rows_by_section")
        sections_complete = section_expected is None or section_expected == sections
        result = {"schema_version":FILE_STAGE_VERSION, "state":"parsed", "source_id":source_id,"format":source["format"],"raw_file":raw_name,"raw_sha256":digest,
                  "registry_sha256":mapping_hash,"staging_sha256":digest_file(temporary/"staging.sqlite3"),
                  "source_contract_sha256":contract_hash,"source_descriptor":source,
                  "original_filename":input_path.name,
                  "counts":counts,"sections":sections,"expected_rows":expected,"received_rows":received,
                  "expected_rows_by_section":section_expected,
                  "complete":expected is not None and expected==received and sections_complete,
                  "mapping_version":source["mapping_version"],"reused":False,
                  "json_number_policy":"integer exact; decimal/exponent literal as text; original bytes authoritative",
                  "quality_status":"synthetic" if source.get("fictional") else "source_unverified"}
        write_json(temporary/"manifest.json",result)
        if output.exists(): raise CollectionError("file_stage_exists_use_new_directory")
        os.rename(temporary,output)
        return result
    except (CollectionError, ValueError, UnicodeError, csv.Error, KeyError, TypeError) as exc:
        # A malformed source is still preserved evidence; never turn it into success0.
        if preserved and str(exc) != "credential_fields_forbidden_in_source_file":
            if c is not None:
                c.rollback()
                c.close()
                c = None
            failure = {"schema_version":FILE_STAGE_VERSION, "state":"parsing_failed", "complete":False,
                       "source_id":source_id, "source_descriptor":source, "source_contract_sha256":contract_hash,
                       "registry_sha256":mapping_hash, "raw_file":raw_name, "raw_sha256":digest,
                       "mapping_version":source["mapping_version"], "staging_rows_committed":0,
                       "error":str(exc) if isinstance(exc, CollectionError) else "file_parsing_failed"}
            write_json(temporary / "manifest.json", failure)
            if not output.exists():
                os.rename(temporary, output)
        raise
    finally:
        if c is not None: c.close()
        if temporary.exists():
            if not temporary.resolve().is_relative_to(output.parent.resolve()) or not temporary.name.startswith(".file-stage-"):
                raise CollectionError("unsafe_temporary_cleanup")
            shutil.rmtree(temporary)


def validate_staged_file(directory: Path, registry_path: Path) -> tuple[dict, dict]:
    """Bind a registered contract to byte-preserved raw and every mapped stage row."""
    manifest = read_json(directory / "manifest.json")
    registry = read_registry(registry_path)
    source = next((s for s in registry["sources"] if s["source_id"] == manifest.get("source_id")), None)
    if not source or source.get("status") != "approved" or not source.get("scope") or not source.get("approval_ref"):
        raise CollectionError("file_source_requires_approved_registry_scope")
    if (manifest.get("schema_version") != FILE_STAGE_VERSION or manifest.get("state") != "parsed"
            or source != manifest.get("source_descriptor")
            or manifest.get("source_contract_sha256") != _content_key(source)
            or manifest.get("mapping_version") != source.get("mapping_version")):
        raise CollectionError("file_stage_contract_mismatch")
    raw_path = safe_file(directory, manifest["raw_file"])
    database = safe_file(directory, "staging.sqlite3")
    if digest_file(raw_path) != manifest["raw_sha256"] or digest_file(database) != manifest["staging_sha256"]:
        raise CollectionError("file_stage_integrity_failed")
    c = sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True)
    c.row_factory = sqlite3.Row
    try:
        if c.execute("PRAGMA integrity_check").fetchone()[0] != "ok" or c.execute("PRAGMA foreign_key_check").fetchall():
            raise CollectionError("file_stage_integrity_failed")
        descriptor = c.execute("SELECT * FROM sources").fetchall()
        if len(descriptor) != 1 or tuple(descriptor[0]) != (source["source_id"], source["format"], source["mapping_version"]):
            raise CollectionError("file_stage_contract_mismatch")
        files = c.execute("SELECT * FROM raw_files").fetchall()
        if (len(files) != 1 or tuple(files[0])[:5] != (manifest["raw_sha256"], source["source_id"],
                manifest["raw_file"], manifest["raw_sha256"], raw_path.stat().st_size)):
            raise CollectionError("file_stage_raw_lineage_mismatch")
        sections = {s: 0 for s in _sections(raw_path, source)}
        counts = {"valid": 0, "rejected": 0}
        for section, position, raw in file_rows(raw_path, source):
            row = c.execute("SELECT * FROM staging_items WHERE file_id=? AND section=? AND row_position=?",
                            (manifest["raw_sha256"], section, position)).fetchone()
            state, reason, standard = "valid", None, None
            try:
                standard = _map_fields(raw, source["field_mapping"])
                if source.get("record_type"):
                    standard = normalize_file_record(standard, source)
            except CollectionError as exc:
                state, reason = "rejected", str(exc)
            expected_id = observation_id_for(source=source["source_id"],
                source_file_ref=_content_key(["section", section]) + "/" + manifest["raw_file"],
                source_content_sha256=manifest["raw_sha256"], source_row_number=position)
            if (row is None or row["observation_id"] != expected_id or json.loads(row["raw_json"]) != raw
                    or (json.loads(row["standard_json"]) if row["standard_json"] is not None else None) != standard
                    or (row["validation_status"], row["reason"]) != (state, reason)):
                raise CollectionError("file_stage_mapped_value_mismatch")
            sections[section] += 1
            counts[state] += 1
        received = sum(counts.values())
        complete = (type(source.get("expected_rows")) is int and source["expected_rows"] == received
                    and (source.get("expected_rows_by_section") is None or source["expected_rows_by_section"] == sections))
        if (c.execute("SELECT COUNT(*) FROM staging_items").fetchone()[0] != received
                or counts != manifest["counts"] or sections != manifest["sections"]
                or manifest["received_rows"] != received or manifest["complete"] != complete):
            raise CollectionError("file_stage_count_reconciliation_failed")
        if not complete:
            raise CollectionError("file_stage_incomplete")
        if counts["rejected"]:
            raise CollectionError("file_stage_rejected_rows_prevent_publication")
        return source, manifest
    finally:
        c.close()
