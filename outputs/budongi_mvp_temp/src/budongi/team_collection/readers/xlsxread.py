"""Read-only XLSX sheet reader yielding provenance-preserving RawRows."""
from __future__ import annotations

import math
from collections.abc import Iterator
from datetime import date, datetime, time
from pathlib import Path

import openpyxl

from budongi.team_collection import contract as C
from budongi.team_collection.base import RawRow


def cell_text(value: object) -> str | None:
    """Convert one openpyxl cell value to its canonical text (None for empty)."""
    if value is None:
        return None
    if isinstance(value, str):
        return value if value != "" else None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isfinite(value) and value == int(value):
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0 and value.microsecond == 0:
            return value.date().isoformat()
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return str(value)


def _parse_header(raw_row: tuple) -> tuple:
    texts = []
    for v in raw_row:
        t = cell_text(v)
        texts.append(t.strip() if t is not None else None)
    width = 0
    for i, t in enumerate(texts):
        if t:
            width = i + 1
    if width == 0:
        raise C.ContractError("invalid_header")
    header = tuple(texts[:width])
    if any(not h for h in header):
        raise C.ContractError("invalid_header")
    if len(set(header)) != len(header):
        raise C.ContractError("invalid_header")
    return header


def _iter_rows(wb, ws, header_row: int, width: int, section: str) -> Iterator[RawRow]:
    try:
        position = header_row
        for raw in ws.iter_rows(min_row=header_row + 1, values_only=True):
            position += 1
            texts = [cell_text(v) for v in raw]
            head = texts[:width]
            head.extend([None] * (width - len(head)))
            tail = texts[width:]
            last = -1
            for i, t in enumerate(tail):
                if t is not None:
                    last = i
            if last >= 0:
                yield RawRow(section, position, tuple(head) + tuple(tail[:last + 1]),
                             "ragged", "ragged_row")
            elif all(t is None for t in head):
                yield RawRow(section, position, tuple(head), "blank", "blank_row")
            else:
                yield RawRow(section, position, tuple(head))
    finally:
        wb.close()


def read_xlsx(path: str | Path, *, sheet: str, header_row: int,
              section: str) -> tuple[tuple, Iterator[RawRow]]:
    """Validate the header row of the named sheet and return (header, lazy RawRows)."""
    if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
        raise C.ContractError("invalid_header_row")
    p = Path(path)
    if not p.is_file():
        raise C.ContractError("file_not_found")
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception:
        raise C.ContractError("invalid_xlsx") from None
    try:
        if sheet not in wb.sheetnames:
            raise C.ContractError("sheet_not_found")
        ws = wb[sheet]
        raw_header = ()
        found = False
        for raw in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
            raw_header = raw
            found = True
            break
        if not found:
            raise C.ContractError("invalid_header")
        header = _parse_header(raw_header)
    except Exception:
        wb.close()
        raise
    return header, _iter_rows(wb, ws, header_row, len(header), section)
