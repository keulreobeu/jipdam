"""Four-sheet metadata workbook writer for role-2 packages."""
from __future__ import annotations

import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableColumn

from budongi.team_collection import contract as C
from budongi.team_collection.base import Cell

_SHEET_NAMES: tuple[str, ...] = (
    "데이터셋 목록",
    "컬럼별 메타데이터",
    "데이터 품질 및 연계",
    "전처리 및 이슈 관리",
)

_HEADERS_RAW: tuple[str, ...] = (
    "데이터셋_ID, 데이터셋명, 분류, 담당자, 제공기관, 출처_URL, 수집방법, 전체행수, 전체컬럼수, 파일수, 파일형식, 인코딩, 데이터기준일, 수집일, 갱신주기, 공간범위, 행단위, PK후보, JOIN키, 좌표계, 저장위치, 가공상태, 이용조건, 활용목적, 중복관계, 측정근거, 검증상태, 원본해시, 수집기록",
    "데이터셋_ID, 컬럼명, 한글명, 설명, 물리자료형, 논리자료형, 전체행수, NULL수, NULL비율, 고유값수, 고유값비율, 최솟값, 최댓값, 평균, 중앙값, 대표값, 단위, 허용값_코드, 이상값수, 이상값비율, 키구분, 전처리필요여부, 의미상결측수, 통합결측수, 통합결측률, 시트_원본위치, 검사대상수, 유효성검사규칙, 측정근거, 정리본_열",
    "데이터셋_ID, 전체중복건수, 전체중복률, PK중복건수, PK중복률, 유효성오류건수, 주요키결측률, 주소보유율, 좌표보유율, 좌표유효율, 공간정보형식, 좌표계, 행정구역기준, 연결대상, JOIN키_방식, JOIN성공률, 미연결건수, 기준시점일치여부, 전체행수, 중복판정기준, 키기준, 주요키결측수, 주소보유수, 좌표보유수, 좌표유효수, 검사대상수, 검사규칙, 검증상태, 근거",
    "이슈_ID, 데이터셋_ID, 대상컬럼, 문제유형, 상세내용, 영향도, 처리방안, 담당자, 처리상태, 처리일, 비고",
)

_EXPECTED_HEADERS: dict[str, tuple[str, ...]] = {
    name: tuple(raw.split(", ")) for name, raw in zip(_SHEET_NAMES, _HEADERS_RAW)
}

_VALID_VALUE_TYPES = (str, int, float, Decimal, bool, datetime.date, datetime.datetime)

_TITLE_FONT = Font(name="Arial", size=15, bold=True, color="253B61")
_NOTE_FONT = Font(name="Arial", size=10, bold=False, color="566273")
_HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="253B61")
_HEADER_ALIGNMENT = Alignment(horizontal="center", wrap_text=True)
_DATA_FONT = Font(name="Arial", size=10, bold=False, color="243247")
_DATA_ALIGNMENT = Alignment(wrap_text=True)


def write_workbook(path: Path, *, sheets: dict[str, dict]) -> dict[str, int]:
    """Write the role-2 four-sheet metadata workbook and return sheet row counts."""
    if isinstance(path, str):
        p = Path(path)
    elif isinstance(path, Path):
        p = path
    else:
        raise C.ContractError("invalid_output_path")

    # R8 (a) path suffix, lower-cased, not '.xlsx' -> 'invalid_output_path'
    if p.suffix.lower() != ".xlsx":
        raise C.ContractError("invalid_output_path")

    # R8 (b) path exists -> 'output_exists'
    if p.exists():
        raise C.ContractError("output_exists")

    # R8 (c) path.parent not an existing directory -> 'output_parent_missing'
    if not p.parent.is_dir():
        raise C.ContractError("output_parent_missing")

    # R8 (d) sheets not a dict or its key set differs from the four names -> 'invalid_sheets'
    if not isinstance(sheets, dict) or set(sheets.keys()) != set(_SHEET_NAMES):
        raise C.ContractError("invalid_sheets")

    # R8 (e) check each sheet in R3 order fully (f..i) before the next
    for name in _SHEET_NAMES:
        sheet_data = sheets[name]

        # R9 (f) sheet value dict structure
        if (
            not isinstance(sheet_data, dict)
            or not isinstance(sheet_data.get("title"), str)
            or not isinstance(sheet_data.get("note"), str)
            or not isinstance(sheet_data.get("headers"), tuple)
            or not isinstance(sheet_data.get("rows"), list)
        ):
            raise C.ContractError(f"invalid_sheet:{name}")

        # R9 (g) headers match R4-R7
        expected_headers = _EXPECTED_HEADERS[name]
        if sheet_data["headers"] != expected_headers:
            raise C.ContractError(f"invalid_headers:{name}")

        # R10 (h) ragged rows and cell type
        expected_len = len(expected_headers)
        rows = sheet_data["rows"]
        for row in rows:
            if not isinstance(row, list) or len(row) != expected_len:
                raise C.ContractError(f"ragged_row:{name}")

        for row in rows:
            for cell in row:
                if not isinstance(cell, Cell):
                    raise C.ContractError(f"invalid_cell:{name}")

        # R11 (i) cell formula and value
        for row in rows:
            for cell in row:
                if cell.formula is not None:
                    if (
                        not isinstance(cell.formula, str)
                        or not cell.formula
                        or cell.formula.startswith("=")
                    ):
                        raise C.ContractError(f"invalid_formula:{name}")
                else:
                    if cell.value is not None and not isinstance(cell.value, _VALID_VALUE_TYPES):
                        raise C.ContractError(f"invalid_cell:{name}")

    # Build workbook
    wb = Workbook()
    # R19 calcPr fullCalcOnLoad="1"
    wb.calcPr = CalcProperties(fullCalcOnLoad=True)

    result: dict[str, int] = {}

    for i, name in enumerate(_SHEET_NAMES, 1):
        if i == 1:
            ws = wb.active
            ws.title = name
        else:
            ws = wb.create_sheet(title=name)

        sheet_data = sheets[name]
        headers = sheet_data["headers"]
        rows = sheet_data["rows"]
        result[name] = len(rows)

        # R12 Layout: A2 title, A3 note
        ws["A2"].value = sheet_data["title"]
        ws["A2"].font = _TITLE_FONT
        ws["A3"].value = sheet_data["note"]
        ws["A3"].font = _NOTE_FONT

        # R13 Header row 5
        for col_idx, header in enumerate(headers, 1):
            hc = ws.cell(row=5, column=col_idx)
            hc.value = header
            hc.font = _HEADER_FONT
            hc.fill = _HEADER_FILL
            hc.alignment = _HEADER_ALIGNMENT

        # R14 Data rows start at row 6
        for k, row_data in enumerate(rows):
            row_idx = 6 + k
            for j, cell_item in enumerate(row_data):
                col_idx = j + 1
                dc = ws.cell(row=row_idx, column=col_idx)
                dc.font = _DATA_FONT
                dc.alignment = _DATA_ALIGNMENT

                # R15 Cell writing
                if cell_item.formula is not None:
                    dc.value = "=" + cell_item.formula
                    dc.data_type = "f"
                elif cell_item.value is not None:
                    dc.value = cell_item.value
                    if isinstance(cell_item.value, str):
                        dc.data_type = "s"

                # R16 Number format
                if cell_item.number_format is not None:
                    dc.number_format = cell_item.number_format
                elif cell_item.formula is None and isinstance(
                    cell_item.value, (datetime.date, datetime.datetime)
                ):
                    dc.number_format = "yyyy-mm-dd"
                else:
                    dc.number_format = "General"

        # R17 Excel table
        last_col = get_column_letter(len(headers))
        last_row = 5 + len(rows) if rows else 6
        table_ref = f"A5:{last_col}{last_row}"
        table_name = f"MetadataTable{i}"
        table_cols = [TableColumn(id=idx, name=header) for idx, header in enumerate(headers, 1)]
        tab = Table(
            id=i,
            name=table_name,
            displayName=table_name,
            ref=table_ref,
            tableColumns=table_cols,
        )
        tab.name = table_name
        tab.displayName = table_name
        tab.tableColumns = table_cols
        ws.add_table(tab)

        # R18 Freeze panes and column width
        ws.freeze_panes = "C6"
        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 18

    wb.save(p)
    return result
