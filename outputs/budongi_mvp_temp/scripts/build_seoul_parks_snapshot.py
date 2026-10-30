"""Build a traceable, unreviewed submission from the approved Seoul park XLSX."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil

from openpyxl import load_workbook

HEADERS = ["연번", "관리부서", "전화번호", "공원명", "공원개요", "면적", "개원일", "주요시설", "주요식물", "안내도", "오시는길", "이용시참고사항", "이미지", "지역", "공원주소", "X좌표(GRS80TM)", "Y좌표(GRS80TM)", "X좌표(WGS84)", "Y좌표(WGS84)", "바로가기"]
SOURCE_ID = "seoul_oa394_major_parks"
BATCH_ID = "seoul-major-parks-2026h2-20261028-v1"
EXPECTED_SHA256 = "9a64b315997327f4d2d60519daeb5b0df044c781f9dc4d6d45626f94f8ecad28"
RETRIEVED = "2026-10-28T00:56:26.793482Z"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def build(source: Path, output: Path) -> None:
    if sha256(source) != EXPECTED_SHA256:
        raise ValueError("source hash does not match the approved download")
    if output.exists():
        raise ValueError("output already exists")
    workbook = load_workbook(source, read_only=True, data_only=False)
    if workbook.sheetnames != ["공원현황"]:
        raise ValueError("unexpected worksheet list")
    sheet = workbook["공원현황"]
    header = [c.value for c in next(sheet.iter_rows(min_row=3, max_row=3, max_col=20))]
    if header != HEADERS:
        raise ValueError("unexpected source headers")
    staged = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    records, missing, names = [], Counter(), Counter()
    for cells in sheet.iter_rows(min_row=4, max_col=20):
        values = [c.value for c in cells]
        if not any(v is not None for v in values):
            continue
        if any(c.data_type == "f" for c in cells):
            raise ValueError(f"formula requires separate review at row {cells[0].row}")
        for i, value in enumerate(values):
            if value is None or value == "":
                missing[HEADERS[i]] += 1
            if isinstance(value, (date, datetime)):
                values[i] = value.isoformat()
        if not isinstance(values[0], int) or not isinstance(values[3], str) or not values[3].strip():
            raise ValueError(f"missing sequence or park name at row {cells[0].row}")
        names[values[3]] += 1
        records.append({"delivery_schema_version": "team_collection_v1", "source_id": SOURCE_ID,
                        "dataset_id": "OA-394", "batch_id": BATCH_ID, "mapping_version": "raw_columns_v1",
                        "record_type": "facility", "entity_level": "facility", "source_entity_id": None,
                        "identifier_namespace": None, "complex_id": None, "building_id": None,
                        "name_raw": values[3], "region_name_raw": values[13], "address_raw": values[14],
                        "latitude": None, "longitude": None, "coordinate_status": "pending",
                        "raw_values": dict(zip(HEADERS, values, strict=True)),
                        "source_file_ref": "raw/seoul-major-parks-2026-h2.xlsx", "source_content_sha256": EXPECTED_SHA256,
                        "source_section": "공원현황", "source_row_number": cells[0].row,
                        "source_retrieved_at_utc": RETRIEVED, "staged_at_utc": staged,
                        "validation_status": "accepted", "validation_reason": None,
                        "quality_status": "source_unverified"})
    workbook.close()
    if len(records) != 155:
        raise ValueError(f"unexpected row count: {len(records)}")
    output.mkdir(parents=True)
    (output / "raw").mkdir()
    (output / "normalized").mkdir()
    (output / "code").mkdir()
    shutil.copyfile(source, output / "raw/seoul-major-parks-2026-h2.xlsx")
    shutil.copyfile(Path(__file__), output / "code/build_seoul_parks_snapshot.py")
    with (output / "normalized/facility.jsonl").open("x", encoding="utf-8", newline="\n") as stream:
        for row in records:
            stream.write(json.dumps(row, ensure_ascii=False, allow_nan=False, separators=(",", ":")) + "\n")
    schema = {"delivery_schema_version": "team_collection_v1", "source_id": SOURCE_ID,
              "record_type": "facility", "mapping_version": "raw_columns_v1", "source_headers": HEADERS,
              "source_section": "공원현황", "source_row_rule": "actual 1-based Excel row, data rows 4-158",
              "identifier": "no official entity identifier asserted; sequence is retained only as raw value",
              "units": "source values only; no area or coordinate conversion",
              "coordinates": "all source coordinates remain in raw_values; latitude/longitude NULL pending review",
              "snapshot_date": None, "source_updated_date": "2026-09-23",
              "source_columns": [{"name": key, "type": "source-preserved", "nullable": True, "unit": None,
                                  "required_when": "source provides a value"} for key in HEADERS],
              "columns": [{"name": key, "type": ("object" if isinstance(value, dict) else
                                                    "integer" if isinstance(value, int) else
                                                    "null" if value is None else "text"),
                           "nullable": key in {"source_entity_id", "identifier_namespace", "complex_id", "building_id",
                                               "region_name_raw", "address_raw", "latitude", "longitude", "validation_reason"},
                           "unit": None, "required_when": "all records"} for key, value in records[0].items()]}
    write_json(output / "schema.json", schema)
    validation = {"raw_rows": 155, "accepted": len(records), "excluded": 0, "rejected": 0,
                  "missing_by_source_column": {h: missing[h] for h in HEADERS},
                  "duplicate_name_groups": {n: c for n, c in names.items() if c > 1},
                  "coordinate_review": "not_run", "coverage": "unknown", "source_quality": "source_unverified",
                  "validated_at_utc": staged, "mapping_version": "raw_columns_v1"}
    write_json(output / "validation.json", validation)
    payload = []
    for relative in ["raw/seoul-major-parks-2026-h2.xlsx", "normalized/facility.jsonl", "schema.json", "validation.json", "code/build_seoul_parks_snapshot.py"]:
        p = output / relative
        payload.append({"path": relative, "bytes": p.stat().st_size, "sha256": sha256(p)})
    manifest = {"delivery_schema_version": "team_collection_v1", "source_id": SOURCE_ID, "dataset_id": "OA-394",
                "provider": "서울특별시", "documentation_url": "https://data.seoul.go.kr/dataList/OA-394/S/1/datasetView.do",
                "batch_id": BATCH_ID, "mapping_version": "raw_columns_v1", "owner": "Codex",
                "approval_ref": "2026-10-28 user approved this single XLSX", "record_url": "https://github.com/keulreobeu/jipdam/issues/40",
                "license_ref": "https://data.seoul.go.kr/dataList/OA-394/S/1/datasetView.do", "license_review_status": "pending",
                "scope": "서울시 주요 공원현황 2026 하반기 XLSX 1건; 전체 공원 커버리지 미확인",
                "source_retrieved_at_utc": RETRIEVED, "provided_at_utc": None, "staged_at_utc": staged,
                "snapshot_date": None, "source_updated_date": "2026-09-23", "quality_status": "source_unverified",
                "requests": [{"collection_status": "complete", "expected_rows": None, "received_rows": 155}],
                "counts": {"raw": 155, "accepted": 155, "excluded": 0, "rejected": 0}, "files": payload}
    write_json(output / "manifest.json", manifest)
    print(json.dumps({"status": "built", "rows": len(records), "source_sha256": EXPECTED_SHA256}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
